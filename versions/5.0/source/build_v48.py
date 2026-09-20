"""由v48_data.json与v48_rules.md生成完整规则书、全卡图鉴与线下版图。
仅离线展示/排版。没有电子游玩逻辑。需要python-docx、beautifulsoup4；PDF导出另用Chromium。
"""
from pathlib import Path
import json, re, html, math, copy
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from bs4 import BeautifulSoup
R=Path(__file__).resolve().parent.parent
W=R/'work'; O=R/'outputs'; O.mkdir(exist_ok=True)
d=json.loads((W/'v48_data.json').read_text(encoding='utf-8'))
base=json.loads((R/'source/v44_data.json').read_text(encoding='utf-8'))
rules=(W/'v48_rules.md').read_text(encoding='utf-8')
E=html.escape
ROM=['初始','I','II','III','IV','V']
KINDS={'chip':'芯片','screen':'屏幕','camera':'影像','body':'机身'}
SANS='Noto Sans CJK SC'; SERIF='Noto Serif CJK SC'
GREEN='244C42'; GOLD='A47E45'; LIGHT='F0F4F0'; INK='263831'
SOURCES=[
 {'id':'S1','maker':'Qualcomm','title':'Snapdragon 778G 5G Mobile Platform','url':'https://www.qualcomm.com/smartphones/products/7-series/snapdragon-778g-5g-mobile-platform','fact':'产品页列出Spectra 570L三ISP。','game':'C02以条件式影像标签匹配表达平台适配；具体匹配规则为桌游设计。'},
 {'id':'S2','maker':'Qualcomm','title':'Snapdragon 870 5G Mobile Platform','url':'https://www.qualcomm.com/smartphones/products/8-series/snapdragon-870-5g-mobile-platform','fact':'产品页将其描述为从骁龙865 Plus演进的平台。','game':'C03的成熟平台减费表达复用方向，不代表真实报价或固定降价幅度。'},
 {'id':'S3','maker':'Qualcomm','title':'Snapdragon 8 Gen 1 Mobile Platform','url':'https://www.qualcomm.com/smartphones/products/8-series/snapdragon-8-gen-1-mobile-platform','fact':'产品页列出18位ISP。','game':'C06降低影像部件的有效科技门槛，是对平台支持能力的抽象，不是现实兼容结论。'},
 {'id':'S4','maker':'MediaTek','title':'MediaTek Dimensity 9300','url':'https://www.mediatek.com/products/smartphones/mediatek-dimensity-9300','fact':'产品页列出全大核CPU设计。','game':'C09提供性能调校路线；额外功耗和吸引力数值不来自真实测量。'},
 {'id':'S5','maker':'Qualcomm','title':'Snapdragon 8 Gen 3 Mobile Platform','url':'https://www.qualcomm.com/smartphones/products/8-series/snapdragon-8-gen-3-mobile-platform','fact':'产品页列出端侧生成式AI能力。','game':'C10以生态标签适配体现平台功能，不表示现实设备会自动获得某种生态服务。'},
 {'id':'S6','maker':'Qualcomm','title':'Snapdragon 8 Elite Mobile Platform','url':'https://www.qualcomm.com/smartphones/products/8-series/snapdragon-8-elite-mobile-platform','fact':'产品页列出Oryon CPU。','game':'C12替代原抽象安全SoC，作为新架构旗舰；节能调校与全部数值为游戏设计。'},
 {'id':'S7','maker':'MediaTek','title':'MediaTek Dimensity 7300','url':'https://www.mediatek.com/products/smartphones/mediatek-dimensity-7300','fact':'产品页强调4nm平台与能效。','game':'C01在后期供货但保留低门槛定位；费用不是实际采购价。'},
 {'id':'S8','maker':'MediaTek','title':'MediaTek Dimensity 8100','url':'https://www.mediatek.com/products/smartphones/mediatek-dimensity-8100','fact':'产品页列出5nm工艺，并强调游戏与能效。','game':'C04保留能效取向；桌游功耗不是瓦数。'},
 {'id':'S9','maker':'MediaTek','title':'MediaTek Dimensity 8200','url':'https://www.mediatek.com/products/smartphones/mediatek-dimensity-8200','fact':'官方产品页定位为4nm级手机平台。','game':'C13调整为低功耗性能平台；调整幅度来自本版数值设计。'},
]
(W/'chip_sources.json').write_text(json.dumps({'inheritedFrom':'4.7','recheckedIn48':False,'sources':SOURCES},ensure_ascii=False,indent=2),encoding='utf-8')
source_md='# 芯片资料与抽象说明\n\n4.8随包资料：沿用4.7引用，4.8未重新核验外部网页。\n\n'+\
'本表区分官方支持的功能与本次创作的游戏效果。没有将厂商营销百分比用作跑分，也没有建立严格自然年供货史。未获得或声称厂商授权、商业背书。各型号名称用于主题辨识，制造费、吸引力、功耗、科技门槛、标签和调校均为桌游参数。\n\n'
for s in SOURCES:source_md+=f"## {s['id']} {s['title']}\n\n来源：{s['maker']}官方产品页。\n\n官方资料：{s['fact']}\n\n游戏抽象：{s['game']}\n\n链接：{s['url']}\n\n"
source_md+='## 未做的验证\n\n未验证真实每一块屏幕、模组、结构件与每颗芯片的量产兼容性；未把不同型号按实际跑分、公平价格或品牌优劣统一排行。部分卡没有特殊能力，仍通过原有数值与标签参与组合。\n'
(O/'芯片资料与抽象说明.md').write_text(source_md,encoding='utf-8')

# ------------------- DOCX complete rulebook -------------------
def set_font(style,size,bold=False,family=SERIF,color=INK):
 style.font.name=family;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor.from_string(color)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),family)
 style.paragraph_format.line_spacing=1.16
 style.paragraph_format.space_after=Pt(5)

def run_font(run,family=None,size=None,color=None,bold=None):
 if family:
  run.font.name=family;run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),family)
 if size:run.font.size=Pt(size)
 if color:run.font.color.rgb=RGBColor.from_string(color)
 if bold is not None:run.bold=bold

def hyperlink(p,label,url=None,anchor=None):
 x=OxmlElement('w:hyperlink')
 if url:
  rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  x.set(qn('r:id'),rid)
 if anchor:x.set(qn('w:anchor'),anchor)
 rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),GREEN);pr.append(co)
 rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),SANS);rf.set(qn('w:hAnsi'),SANS);rf.set(qn('w:eastAsia'),SANS);pr.append(rf)
 rr.append(pr);tt=OxmlElement('w:t');tt.text=label;rr.append(tt);x.append(rr);p._p.append(x)

def rulebook():
 doc=Document(); sec=doc.sections[0]
 sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.top_margin=Mm(19);sec.bottom_margin=Mm(17);sec.left_margin=Mm(18);sec.right_margin=Mm(18)
 sec.header_distance=Mm(8);sec.footer_distance=Mm(8)
 sec.different_first_page_header_footer=True
 for name,size,bold,fam,col in [('Normal',10.5,False,SERIF,INK),('Title',44,True,SANS,GREEN),('Subtitle',15,False,SANS,GOLD),('Heading 1',16,True,SANS,GREEN),('Heading 2',11.5,True,SANS,GREEN)]:
  set_font(doc.styles[name],size,bold,fam,col)
 for name in ['Heading 1','Heading 2']:
  s=doc.styles[name];s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 8);s.paragraph_format.space_after=Pt(5)
 for st in doc.styles:
  for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 doc.styles['Subtitle'].font.italic=False
 doc.styles['Normal'].paragraph_format.widow_control=True
 hp=sec.header.paragraphs[0];hp.text='旗舰元年  /  完整规则书';run_font(hp.runs[0],SANS,8,GREEN)
 hp.paragraph_format.tab_stops.add_tab_stop(Mm(174),WD_ALIGN_PARAGRAPH.RIGHT)
 r=hp.add_run('\t4.8  综合测试版');run_font(r,SANS,8,GOLD)
 fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run_font(fp.add_run('旗舰元年   ·   '),SANS,8,GREEN)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)
 doc.core_properties.title='旗舰元年 4.8 完整规则书';doc.core_properties.subject='稀缺科技、定价收入与无纸笔组件；从4.7延续'
 doc.core_properties.author='';doc.core_properties.keywords='桌游,旗舰元年,4.8,测试版'
 settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
 def p(txt='',style=None):return doc.add_paragraph(txt,style)
 def tab(headers,rows,widths=None,fontsize=9.2):
  n=len(headers); width=174
  widths=widths or ([31,143] if n==2 else [38,58,78] if n==3 else [38,48,42,46] if n==4 else [32,31,37,37,37])
  compact = headers[0]=='定价'
  if compact and doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
  t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  for c,w in zip(t.columns,widths):c.width=Mm(w)
  bp=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','left','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CFDAD2');bp.append(el)
  t._tbl.tblPr.append(bp)
  for ri,vals in enumerate([headers]+rows):
   row=t.rows[0] if ri==0 else t.add_row();rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
   if ri==0:rp.append(OxmlElement('w:tblHeader'))
   for cell,txt,w in zip(row.cells,vals,widths):
    cell.width=Mm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),GREEN if ri==0 else ('F2F5F1' if ri%2==0 else 'FFFFFF'));cp.append(sh)
    mar=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     a=OxmlElement('w:'+edge);a.set(qn('w:w'),'85');a.set(qn('w:type'),'dxa');mar.append(a)
    cp.append(mar);q=cell.paragraphs[0];q.paragraph_format.line_spacing=1.1;q.paragraph_format.space_before=Pt(1.5);q.paragraph_format.space_after=Pt(1.5)
    if ri==0 or (compact and ri<len(rows)):q.paragraph_format.keep_with_next=True
    run_font(q.add_run(str(txt)),SANS if ri==0 else SERIF,fontsize,'FFFFFF' if ri==0 else INK,ri==0)
  q=p();q.paragraph_format.space_after=Pt(2);q.paragraph_format.line_spacing=Pt(2);q.paragraph_format.space_before=Pt(0)
  return t
 # Cover: typography first, no random artwork.
 q=p('FLAGSHIP ERA   /   4.8');q.paragraph_format.space_before=Pt(34);run_font(q.runs[0],SANS,12,GOLD,True)
 q=p('旗舰元年','Title');q.paragraph_format.space_before=Pt(32);q.paragraph_format.space_after=Pt(12)
 p('完整规则书','Subtitle')
 q=p('稀缺科技  ·  收入定价  ·  实体选择');q.paragraph_format.space_before=Pt(18);run_font(q.runs[0],SANS,13,GREEN)
 q=p('2—4人     90—150分钟（待实测）');q.paragraph_format.space_before=Pt(28);run_font(q.runs[0],SANS,10,GOLD)
 q=p('公开抢到的零件，只是别人看得见的一半。\n真正的旗舰，在共同发布时才揭晓。');q.paragraph_format.space_before=Pt(20);q.paragraph_format.space_after=Pt(22);run_font(q.runs[0],SERIF,16,GREEN)
 tab(['保留骨架','本次整合','配套图鉴'],[['五行动强度\n双轨与共同发布','科技稀缺化＋定价收入\n从4.7继续迭代','280种内容牌\n按份数准备352张']], [58,58,58],10)
 q=p('综合测试版  /  2026年9月20日');q.paragraph_format.space_before=Pt(20);run_font(q.runs[0],SANS,10,GOLD)
 q=p('科技0—20，1点折算5；世代2／4／6／8。\n定价牌与调校标记替代纸笔。尚待多人平衡实测。');run_font(q.runs[0],SANS,9,'65766A')
 doc.add_page_break()
 # Contents; field references render actual page numbers in LO/Word.
 p('阅读导航','Heading 1')
 p('初次游玩先读第一至第十三节；遇到结算疑问查第十四节。全卡数值与完整效果见配套图鉴，第十六节提供编号和供货批次索引。')
 headings=re.findall(r'^## (.+)$',rules,re.M)
 for idx,h in enumerate(headings,1):
  q=p();q.paragraph_format.space_after=Pt(6);q.paragraph_format.tab_stops.add_tab_stop(Mm(169),WD_ALIGN_PARAGRAPH.RIGHT)
  hyperlink(q,h,anchor=f'ch_{idx}')
  q.add_run('\t');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),f'PAGEREF ch_{idx} \\h');q._p.append(f)
 p('三条容易混淆的界线','Heading 2')
 p('供应世代 ≠ 个人科技门槛；配件标签 ≠ 公司标签；销售现金 ≠ 收入轨。')
 p('更新说明','Heading 2')
 p('旧版未改动规则已融入正文，不需要边玩边翻旧版补丁。所有数值为测试参数，未引入电子游玩功能。')
 doc.add_page_break()
 # Intro from source MD excluding repeated title/subtitles.
 pre=rules.split('## 一 ')[0].splitlines()
 for ln in pre[4:]:
  if ln.strip():p(ln)
 lines=rules[rules.index('## 一 '):].splitlines();i=0;chapter=0
 while i<len(lines):
  ln=lines[i].strip()
  if not ln:i+=1;continue
  if ln.startswith('## '):
   chapter+=1;q=p(ln[3:],'Heading 1')
   a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(chapter));a.set(qn('w:name'),f'ch_{chapter}');q._p.insert(0,a)
   z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(chapter));q._p.append(z)
  elif ln.startswith('### '):p(ln[4:],'Heading 2')
  elif ln.startswith('|'):
   arr=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',v) for v in cells):arr.append(cells)
    i+=1
   n=len(arr[0]);widths=None
   if n==4 and '路线' in arr[0][0]:widths=[32,54,44,44]
   if n==4 and '世代' in arr[0][0]:widths=[22,92,30,30]
   if n==5:widths=[25,37,37,37,38]
   if n==2 and arr[0][0]=='类别':widths=[30,144]
   tab(arr[0],arr[1:],widths);continue
  else:p(ln)
  i+=1
 p('沿用4.7的官方资料索引','Heading 2')
 p('以下索引及描述沿用4.7存档，本轮未重新核验网页。游戏参数为桌游设计，详见随包资料存档。').paragraph_format.keep_with_next=True
 for s in SOURCES:
  q=p();q.paragraph_format.space_after=Pt(5)
  hyperlink(q,f"[{s['id']}] {s['maker']} · {s['title']}",url=s['url'])
  run_font(q.add_run('  '+s['fact']),SERIF,9.5)
 doc.save(O/'旗舰元年_4.8_完整规则书.docx')

# ------------------- catalog: every design, complete text -------------------
CAT_CSS='''
*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#dfe5de;color:#243c34;font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif}a{color:inherit}header.tools{position:sticky;top:0;z-index:5;background:#244c42;color:#f6f4e8;padding:13px 22px;display:flex;align-items:center;gap:16px;box-shadow:0 2px 12px #17372d35}header.tools b{font-size:16px}header.tools a,header.tools button{background:transparent;color:inherit;border:1px solid #96ad9e88;padding:5px 10px;text-decoration:none;font:12px inherit}header.tools .links{display:flex;gap:7px;flex-wrap:wrap;flex:1}.sheet{width:210mm;height:296mm;margin:16px auto;background:#fffef8;padding:10mm;position:relative;box-shadow:0 4px 20px #17372d24;overflow:hidden;break-after:page;page-break-after:always}.sheet:last-child{break-after:auto;page-break-after:auto}.page-head{height:13mm;display:flex;align-items:flex-start;justify-content:space-between;border-bottom:.4mm solid #244c42;margin-bottom:4mm}.page-head h2{font-size:16px;font-weight:700;letter-spacing:1px;margin:0}.page-head span{font-size:10px;color:#758373;margin-top:4px}.grid{display:grid;grid-template-columns:repeat(2,1fr);grid-template-rows:repeat(3,82mm);gap:4mm}.card{--accent:#416c58;border:.3mm solid #b7c7bb;border-top:1.25mm solid var(--accent);border-radius:2.2mm;background:#fff;position:relative;height:82mm;padding:3.5mm;display:flex;flex-direction:column;overflow:hidden;box-shadow:0 1mm 2mm #254a3710}.card.phone{--accent:#416c58}.card.tech{--accent:#397685}.card.business{--accent:#977246}.card.common{--accent:#606958}.card.custom{--accent:#8c6478}.card.starter{--accent:#747a64}.card.project{--accent:#7e743c}.card.goal{--accent:#686486}.card.action{--accent:#344e5d}.card-top{display:flex;align-items:center;justify-content:space-between;font-size:9px;color:var(--accent);letter-spacing:.4px;height:4mm;flex-shrink:0}.card-top b{font-size:12px;letter-spacing:1px}.card h3{font-family:"Noto Serif CJK SC",SimSun,serif;font-size:16px;line-height:1.35;margin:1.6mm 0 1.8mm;letter-spacing:.3px;font-weight:700;color:#233e34}.card h3.long{font-size:14px}.subline{font-size:10px;color:#697c6a;margin-bottom:2mm;line-height:1.5}.tag{display:inline-block;border:1px solid #ccd7cb;padding:0 1.3mm;margin-right:1mm;border-radius:1mm;font-size:9px;background:#f5f7f2;white-space:nowrap}.metrics{display:grid;grid-template-columns:repeat(3,1fr);gap:1mm;background:#eff4ef;padding:1.6mm 1mm;margin-bottom:2mm;flex-shrink:0}.metric{text-align:center;font-size:9px;color:#647669;line-height:1.4}.metric b{display:block;color:#244c42;font-size:17px;line-height:1.35;font-family:Georgia,"Noto Sans CJK SC",serif}.gate{padding:1.5mm 2mm;background:#f7f0e1;border-left:.7mm solid #b99965;color:#685437;font-size:9px;line-height:1.55;margin-bottom:2mm}.effect{font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;font-size:10px;line-height:1.65;color:#253d34;margin:0}.effect strong{font-weight:700;color:var(--accent)}.effect.dense{font-size:9.5px;line-height:1.56}.life{font-size:8px;line-height:1.45;color:#6f7c6b;border-top:1px solid #dae1d6;margin-top:auto;padding-top:1.8mm}.parttable{width:100%;border-collapse:collapse;table-layout:fixed;font-size:9px;margin:0 0 2mm;text-align:center}.parttable th{background:#eff3ee;color:#596e5b;font-weight:500;padding:1.1mm .4mm;font-size:8px}.parttable td{padding:1.25mm .4mm;border-bottom:1px solid #e1e7df;line-height:1.4}.parttable td:first-child{text-align:left;width:25%;font-weight:600}.parttable td small{display:block;color:#819078;font-weight:400;font-size:7px}.parttable b{font-size:11px;font-weight:500}.tiergrid{display:grid;grid-template-columns:repeat(3,1fr);gap:2mm;margin:3mm 0}.tier{border:1px solid #d9dccc;background:#f8f7ee;text-align:center;padding:2.5mm .5mm;font-size:10px;line-height:1.8}.tier b{font-size:19px;display:block;font-family:Georgia,serif}.tier small{display:block;color:#7e805f;font-size:8px}.goalmark{font-size:26px;font-family:Georgia,serif;color:var(--accent);background:#f3f1f8;border:1px solid #dcd9e7;padding:3mm;margin:3mm 0;text-align:center}.goalmark small{font-family:inherit;font-size:10px}.page-foot{position:absolute;left:10mm;right:10mm;bottom:7mm;border-top:1px solid #bdcab9;display:flex;justify-content:space-between;padding-top:2mm;font-size:8px;color:#70816c}.cover .eyebrow{letter-spacing:2px;font-size:11px;color:#9c7947;margin-top:14mm}.cover h1{font:700 50px "Noto Serif CJK SC",SimSun,serif;margin:14mm 0 3mm;color:#234d41}.cover .subtitle{font-size:23px;letter-spacing:4px;margin-bottom:9mm}.cover .tagline{font-family:"Noto Serif CJK SC",SimSun,serif;font-size:16px;line-height:2;margin:9mm 0}.cover .counts{display:grid;grid-template-columns:repeat(3,1fr);border-top:1px solid #b4c4b5;border-bottom:1px solid #b4c4b5;padding:7mm 0;margin:10mm 0}.counts div{text-align:center;font-size:12px;color:#6b806b}.counts b{font:40px Georgia,serif;display:block;color:#244c42;margin-bottom:3mm}.cover .panel{background:#eff3eb;border-left:1mm solid #a6844b;padding:6mm;line-height:1.9;font-size:12px;margin:9mm 0}.cover small{font-size:10px;color:#72816d;line-height:1.9}.legend h3{font-size:15px;color:#244c42;border-bottom:1px solid #c5d0c0;padding-bottom:2mm;margin:6mm 0 3mm}.legend p{font-size:11px;line-height:1.9;margin:2mm 0}.legend table{border-collapse:collapse;width:100%;font-size:11px}.legend th,.legend td{border-bottom:1px solid #d4dfd0;padding:2mm;text-align:left}.legend th{background:#edf3eb}.aid h3{font-size:15px;color:#244c42}.aid p{font-size:11px;line-height:1.9}.aid table{width:100%;border-collapse:collapse;margin:3mm 0 6mm;font-size:12px}.aid td,.aid th{border:1px solid #c1cebd;padding:3mm;text-align:center}.aid th{background:#edf3e9}.record{border:1px solid #b5c5b3;padding:5mm;margin-top:5mm;position:relative}.record h4{font-size:15px;margin:0 0 3mm}.record p{margin:2mm 0}.source-item{font-size:9.5px;line-height:1.7;margin-bottom:3mm}.source-item a{font-weight:600;text-decoration:underline}.muted{color:#74836f}@page{size:A4 portrait;margin:0}@media print{body{background:white}header.tools{display:none}.sheet{margin:0;box-shadow:none}*{-webkit-print-color-adjust:exact;print-color-adjust:exact}}@media screen and (max-width:820px){header.tools{position:static;flex-wrap:wrap}.sheet{margin:10px auto;transform-origin:top left}}
'''

CAT_CSS+='''\n.effect{font-size:12px;line-height:1.6}.effect.dense{font-size:11.5px;line-height:1.55}.life{font-size:9px;line-height:1.45}.subline{font-size:11px}.card-top{font-size:10px}.metric{font-size:10px}.phone .metrics{grid-template-columns:repeat(5,1fr)}.gate{font-size:10px}.parttable{font-size:10px}.parttable th{font-size:9px}.parttable td{padding:1.1mm .4mm}.parttable td small{font-size:8px}.parttable b{font-size:12px}.tag{font-size:10px}\n'''

def tags(arr):return ''.join(f'<span class="tag">{E(x)}</span>' for x in arr)
def metric(vals):return '<div class="metrics">'+''.join(f'<div class="metric">{E(k)}<b>{E(str(v))}</b></div>' for k,v in vals)+'</div>'
def card(c,typ):
 cid=c['id'];qty=c.get('qty',1)
 titles={'phone':'产品企划','tech':'科技','business':'企业','common':'稳定配件','custom':'定制配件','starter':'初始稳定','project':'产业突破','goal':'终局目标'}
 if typ in ['common','custom']:badge=f"{ROM[c['era']]}代 · {qty}张"
 else:badge=f'{titles[typ]} · {qty}张'
 s=f'<article class="card {typ}" id="card-{cid}" data-card-id="{cid}" data-card-type="{typ}"><div class="card-top"><b>{cid}</b><span>{badge}</span></div>'
 cls='long' if len(c['name'])>11 else ''
 s+=f'<h3 class="{cls}">{E(c["name"])}</h3>'
 life=''
 if typ=='phone':
  s+=f'<div class="subline">{E(c["market"])}　{tags(c["tags"])}</div>'
  s+=metric([('组装',c['strength']),('科技门槛',c['tech']),('开发费',c['development']),('验收',c['appeal']),('收入',c['income'])])
  gate=' 且 '.join(f'{k}{v}' for k,v in c.get('tagRequirements',{}).items()) or '无'
  s+=f'<div class="gate">公司标签门槛：{E(gate)}</div>'
  s+=f'<p class="effect">{E(c["effect"])}</p>'
  life='标签门槛只看发布前公司。验收与开发费不是吸引力或采购价。上市后保留，不重复结算。'
 elif typ in ['tech','business']:
  s+=f'<div class="subline">{tags(c["tags"])}</div>'
  vals=[('行动要求',c['strength']),('科技门槛',c['tech']),('费用',c['cost'])]
  s+=metric(vals)
  if typ=='tech':
   s+=f'<div class="gate">{c["researchClass"]} · '+(f'打出＋1科技<br>研发基础：此前已打出至少{c["priorTechCardsRequired"]}张科技卡' if c['gain'] else '打出不增加科技；持续能力照常生效')+'</div>'
  cls='dense' if len(c['effect'])>95 else ''
  s+=f'<p class="effect {cls}">{E(c["effect"])}</p>'
  life='通过研发打出。本卡不计入自身研发基础；科技收益以印刷值为准。持续保留。' if typ=='tech' else '通过合作打出并持续保留。每场发布型效果每张每场最多一次；费用减免仍受最低支付限制。'
 elif typ in ['common','custom','starter']:
  cs=c['components'];kind=KINDS[cs[0]['kind']] if len(cs)==1 else f'{len(cs)}项整合配件'
  s+=f'<div class="subline">{kind}　'+('初始分发，不混牌库' if typ=='starter' else '公开供应 / 永久入库' if typ=='common' else '秘密手牌 / 一次性')+'</div>'
  s+='<table class="parttable"><colgroup><col style="width:22%"><col style="width:13%"><col style="width:13%"><col style="width:16%"><col style="width:13%"><col style="width:23%"></colgroup><thead><tr><th>部件</th><th>制造</th><th>吸引力</th><th>功／供</th><th>科技</th><th>标签</th></tr></thead><tbody>'
  for v in cs:
   form='折叠' if v['form']=='fold' else ('直板' if v['kind'] in ['screen','body'] else '通用结构')
   power=('供' if v['kind']=='body' else '耗')+str(v['load'])
   label=(KINDS[v['kind']]+('·'+form if v['kind'] in ['screen','body'] else '')) if len(cs)>=3 else KINDS[v['kind']]+f'<small>{form}</small>'
   s+=f'<tr><td>{label}</td><td><b>{v["cost"]}</b></td><td><b>{v["appeal"]}</b></td><td>{power}</td><td>{v["tech"]}</td><td>{E(v["tag"])}</td></tr>'
  s+='</tbody></table>'
  name=f'<strong>{E(c["abilityName"])}。</strong>' if c.get('abilityName') else ''
  cls='dense' if len(c['effect'])>78 or len(cs)>=3 else ''
  s+=f'<p class="effect {cls}">{name}{E(c["effect"])}</p>'
  life=c['lifeCycle']
 elif typ=='project':
  label=c['label']+('公司标签数' if c['metric']=='tag' else '')
  # "性能公司标签数" keeps company/part distinction explicit.
  s+=f'<div class="subline">统计：{E(label)}</div><div class="tiergrid">'
  for ii,(th,reward) in enumerate(zip(c['thresholds'],c['rewards']),1):s+=f'<div class="tier"><small>第{ii}档 · 要求{th}</small><b>＋{reward}</b>科技</div>'
  s+='</div><p class="effect">'+E(c['effect'])+'</p>'
  life='普通合作需强度4，升级需强度3。每局随机展示4个项目；每档全桌一人，首次完成达成A4。'
 elif typ=='goal':
  s+='<div class="goalmark">6 <small>分上限</small></div>'
  s+=f'<p class="effect">{E(c["effect"])}</p>'
  life='开局抽2留1，秘密保存。只在最终发布及成长之后计分，不参与提前触发终局。'
 s+=f'<div class="life">{E(life)}</div></article>'
 return s

def catalog():
 groups=[('phone','产品企划',[c for c in d['cards'] if c['type']=='phone']),('tech','科技卡',[c for c in d['cards'] if c['type']=='tech']),('business','企业卡',[c for c in d['cards'] if c['type']=='business']),('common','稳定配件',sorted([c for c in d['parts'] if c['type']=='common'],key=lambda c:('CDIF'.index(c['id'][0]),c['id']))),('custom','定制配件',[c for c in d['parts'] if c['type']=='custom']),('starter','初始稳定配件',d['starters']),('project','产业突破',d['projects']),('goal','终局目标',d['goals'])]
 page=3;index=[]
 for typ,title,cs in groups:index.append((typ,title,page,len(cs),sum(c.get('qty',1) for c in cs)));page+=math.ceil(len(cs)/6)
 index.extend([('action','双面行动卡',page,5,'每人5张'),('aid','定价与科技速查',page+2,'—','不计内容牌')]);total=page+3
 sheets=[]
 def sheet(body,title=None,num=None,cls='',anchor=''):
  num=num or len(sheets)+1
  head='' if title is None else f'<div class="page-head"><h2>{E(title)}</h2><span>旗舰元年 4.8 · 全卡牌图鉴</span></div>'
  sheets.append(f'<section class="sheet {cls}" id="{anchor or "page-"+str(num)}">{head}{body}<div class="page-foot"><span>4.8 综合测试版 · 卡牌数值以本版为准</span><span>{num:02d} / {total:02d}</span></div></section>')
 cover='''<div class="eyebrow">FLAGSHIP ERA / DESIGN CATALOGUE</div><h1>旗舰元年</h1><div class="subtitle">全卡牌图鉴</div><div class="tagline">同一套供应链，\n可以造出不一样的旗舰。</div><div class="counts"><div><b>280</b>种内容牌设计</div><div><b>352</b>张内容牌实体</div><div><b>4.8</b>综合测试版</div></div><div class="panel">完整收录企划、科技、企业、稳定配件、定制配件、初始配件、产业突破与终局目标。相同编号只展示一次，按标注份数准备。<br>另附五种行动卡的两面及无纸笔定价、调校说明。</div><small>从4.7延续，全部编号保留；实施科技稀缺化与实体定价。<br>科技0—20，每点折算5；供应解锁2／4／6／8。定价直接影响收入和现金。<br>内容牌为完整参考排版，尚非量产刀模；定价牌正反面和标记另附实物打印稿。<br>2026年9月20日 · 尚待充分多人实测</small>'''
 sheet(cover,cls='cover',anchor='catalog-cover')
 legend='<h3>分类与页码</h3><table><tr><th>分类</th><th>设计／实体</th><th>起始页</th></tr>'
 for typ,title,pg,designs,qty in index:legend+=f'<tr><td><a href="#cat-{typ}">{title}</a></td><td>{designs}／{qty}</td><td>{pg}</td></tr>'
 legend+='</table><h3>读懂卡面</h3><p>企划的“验收”是要求，不是企划额外提供的吸引力。“公司标签门槛”只看发布前已经公开的公司卡；配件标签不计公司。</p><p>部件“制造”是每次使用的制造费，不是一次采购价。“功／供”中，耗代表芯片、屏幕、影像功耗，供代表机身供电。多项定制的每行都必须装入同一台。</p><p>卡面世代决定何时供货；科技门槛决定你能否在发布时使用。调校在组装或返工时选，定价在发布揭晓前同时选。印刷数值不因这些修正而改变。</p><h3>完整玩法与实物准备</h3><p>先阅读配套《4.8完整规则书》。图鉴中的六卡一页布局用于检索，不要求把所有复制份数重复印在图鉴中。行动卡每位玩家另备一套，不属于352张内容牌。</p>'
 sheet(legend,'阅读图鉴',cls='legend',anchor='catalog-index')
 manifest=[]
 for typ,title,cs in groups:
  for off in range(0,len(cs),6):
   subset=cs[off:off+6];pg=len(sheets)+1
   body='<div class="grid">'+''.join(card(c,typ) for c in subset)+'</div>'
   sheet(body,f'{title}  /  {subset[0]["id"]}—{subset[-1]["id"]}',cls='cards-page',anchor=f'cat-{typ}' if off==0 else '')
   for c in subset:manifest.append({'id':c['id'],'name':c['name'],'type':typ,'qty':c.get('qty',1),'page':pg})
 acts=[]
 for a in d['actions']:
  for side,label in [('front','普通面'),('back','升级面')]:
   acts.append(f'<article class="card action" data-action-id="{a["id"]}-{side}"><div class="card-top"><b>{a["id"]}</b><span>{label} · 每人1张双面卡</span></div><h3>{a["name"]}</h3><div class="gate">强度S＝当前所在位置（1—5）</div><p class="effect">{E(a[side])}</p><div class="life">用后移回强度1，其左侧行动依次右移。可改整备：只获3现金，不触发原行动。详细边界按完整规则书。</div></article>')
 for off in range(0,len(acts),6):sheet('<div class="grid">'+''.join(acts[off:off+6])+'</div>','双面行动卡 / 普通面与升级面',cls='cards-page',anchor='cat-action' if off==0 else '')
 aid='<h3>市场收入与现金</h3><p>格内数字为“收入／现金”。每公司每市场只由代表机领一次；不再额外加旧版3／1。</p><table><tr><th>定价</th><th>吸引力</th><th>第一</th><th>第二</th><th>第三／第四</th></tr>'
 for p in d['pricing']['strategies']:
  aid+=f'<tr><td>{p["name"]}</td><td>{p["appealDelta"]:+d}</td>'+''.join(f'<td>{p["incomeByRank"][i]}／{p["cashByRank"][i]}</td>' for i in range(3))+'</tr>'
 aid+='</table><h3>每人9张定价牌 · 扣牌锁定后同时翻开</h3><p>筹备位1、2、3各有走量、标准、溢价三张，每个已筹备位置选择对应编号的一张扣下，其他牌留在屏后。允许所有机型选同一价格。漏放、多放或编号不符按标准价。结算后收回，不占手牌、不消耗、不留在已上市企划下面。</p><h3>每人3枚调校标记 · 不写编号</h3><p>组装／返工时，标记放在该机提供方案的配件卡上；选T19时放对应筹备位的T19节能格。没有标记就是默认。只有正常返工可重选；揭晓时按标记位置读取，不能临时改模式。</p><h3>稀缺科技速查</h3><table><tr><th>来源</th><th>科技</th><th>限制</th></tr><tr><td>24张应用技术</td><td>0</td><td>仍给能力、公司标签</td></tr><tr><td>T07—T18前沿技术</td><td>各1</td><td>此前已有2张科技卡</td></tr><tr><td>产业突破</td><td>1／2／3</td><td>每人每项目一次；全桌争档</td></tr><tr><td>基础研究</td><td>1</td><td>每人整局一次；普通S4付8，升级S3付6</td></tr></table><p>科技0—20，每点折算5；供应世代2／4／6／8；科技2、5各升级一次；科技3且上市2张解锁第三筹备位。仅8张指定企划保留科技奖励，完整条件见卡面。</p><p class="muted">实物定价牌正反面、T19小垫板与标记见《4.8实体定价与标记》。按100%尺寸打印，不需要纸笔、擦写笔或手机记录。</p>'
 sheet(aid,'定价、科技与实物速查',cls='aid',anchor='cat-aid')
 src='<p>官方资料只作为功能方向参考。所有费用、吸引力、功耗、门槛、标签与调校数值均为本次桌游设计，非实测或厂商承诺。供货世代不等于严格自然年。</p>'
 for s in SOURCES:src+=f'<div class="source-item"><a href="{E(s["url"])}">[{s["id"]}] {E(s["maker"])} · {E(s["title"])}</a><br>{E(s["fact"])} {E(s["game"])}</div>'
 src+='<h3>版本与数据追溯</h3><p>底稿：4.7完整交付数据与规则。4.8调整全部科技门槛、科技卡收益与部分费用、25张企划奖励、O11门槛及定价收益。原编号保留，请勿混用4.7科技门槛或收益。</p><p>随包提供v48_data.json、完整规则Markdown、逐字段差异、生成脚本与定向校验结果。程序校验不构成多人策略或时长实测。</p><p>芯片功能资料沿用4.7存档；4.8未重新核验外部网页。此处保留来源追溯，不是本轮最新硬件研究。</p>'
 sheet(src,'资料与版本说明',cls='legend',anchor='catalog-sources')
 assert len(sheets)==total,(len(sheets),total)
 links=''.join(f'<a href="#cat-{t}">{name}</a>' for t,name,_,_,_ in index)
 h=f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>旗舰元年 4.8 · 全卡牌图鉴</title><style>{CAT_CSS}</style></head><body><header class="tools"><b>旗舰元年 4.8</b><div class="links">{links}</div><button onclick="window.print()">打印图鉴</button></header>'+''.join(sheets)+'</body></html>'
 (O/'旗舰元年_4.8_全卡牌图鉴.html').write_text(h,encoding='utf-8')
 (W/'catalog_manifest.json').write_text(json.dumps({'pages':total,'cards':manifest},ensure_ascii=False,indent=2),encoding='utf-8')
 # Plain-text full catalogue, useful for future editing, not a truncated subset.
 md='# 旗舰元年4.8全卡牌条目\n\n280种、352张内容牌；另附5种行动卡的两面。以本版完整规则裁定时点。\n\n'
 for typ,title,cs in groups:
  md+=f'## {title}\n\n'
  for c in cs:
   md+=f'### {c["id"]} {c["name"]}（{c.get("qty",1)}张）\n\n'
   if typ in ['common','custom','starter']:
    md+=f'供应：{ROM[c.get("era",0)]}；'+('公开稳定' if typ!='custom' else '主牌定制')+'。\n\n'
    for v in c['components']:md+=f'- {KINDS[v["kind"]]}：{v["name"]}；制造{v["cost"]}，印刷吸引力{v["appeal"]}，'+('供电' if v['kind']=='body' else '功耗')+f'{v["load"]}，科技{v["tech"]}，标签{v["tag"]}，形态{v["form"]}。\n'
    md+='\n'+c['effect']+'\n\n'+c['lifeCycle']+'\n\n'
   elif typ=='phone':
    md+=f'市场：{c["market"]}；企划标签：'+ '、'.join(c['tags'])+f'；组装{c["strength"]}；科技{c["tech"]}；开发费{c["development"]}；验收{c["appeal"]}；收入{c["income"]}。\n\n公司标签门槛：'+('且'.join(k+str(v) for k,v in c.get('tagRequirements',{}).items()) or '无')+'。\n\n'+c['effect']+'\n\n'
   elif typ in ['tech','business']:
    md+=f'行动{c["strength"]}；科技门槛{c["tech"]}；费用{c["cost"]}；标签'+ '、'.join(c['tags'])+('；'+c['researchClass']+'；获得科技'+str(c['gain'])+'；研发基础：此前'+str(c['priorTechCardsRequired'])+'张科技卡' if typ=='tech' else '')+'。\n\n'+c['effect']+'\n\n'
   elif typ=='project':md+='统计：'+c['label']+'；门槛'+str(c['thresholds'])+'；科技奖励'+str(c['rewards'])+'。\n\n'+c['effect']+'\n\n'
   else:md+=c['effect']+'\n\n'
 md+='## 行动卡\n\n'
 for a in d['actions']:md+=f'### {a["id"]} {a["name"]}\n\n普通面：{a["front"]}\n\n升级面：{a["back"]}\n\n'
 (W/'v48_catalog.md').write_text(md,encoding='utf-8')
 print('Catalog:',len(manifest),'designs;',sum(v['qty'] for v in manifest),'physical cards;',total,'pages')


from boards_and_aids_v48 import build_boards, build_aids

if __name__=="__main__":
 rulebook(); catalog(); build_boards(); build_aids()
 print("Built all 4.8 DOCX and HTML outputs.")
