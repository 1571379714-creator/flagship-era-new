"""生成5.0完整规则书、全部卡牌图鉴；不含电子游玩功能。"""
from v50_common import *
import re,math,copy
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
rules=(W/'v50_rules.md').read_text(encoding='utf8')
SANS='Noto Sans CJK SC';SERIF='Noto Serif CJK SC'
GREEN='244C42';GOLD='A47E45';LIGHT='F0F4F0';INK='263831'
def set_font(style,size,bold=False,family=SERIF,color=INK):
 style.font.name=family;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor.from_string(color)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),family)
 style.paragraph_format.line_spacing=1.16
 style.paragraph_format.space_after=Pt(5)

def run_font(run,family=None,size=None,color=None,bold=None):
 if family:
  run.font.name=family;run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),family)
 if size:run.font.size=Pt(size)
 if color:run.font.color.rgb=RGBColor.from_string(color)
 if bold is not None:run.bold=bold

def hyperlink(p,label,url=None,anchor=None):
 x=OxmlElement('w:hyperlink')
 if url:
  rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  x.set(qn('r:id'),rid)
 if anchor:x.set(qn('w:anchor'),anchor)
 rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),GREEN);pr.append(co)
 rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),SANS);rf.set(qn('w:hAnsi'),SANS);rf.set(qn('w:eastAsia'),SANS);pr.append(rf)
 rr.append(pr);tt=OxmlElement('w:t');tt.text=label;rr.append(tt);x.append(rr);p._p.append(x)

def rulebook():
 doc=Document(); sec=doc.sections[0]
 sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.top_margin=Mm(19);sec.bottom_margin=Mm(17);sec.left_margin=Mm(18);sec.right_margin=Mm(18)
 sec.header_distance=Mm(8);sec.footer_distance=Mm(8)
 sec.different_first_page_header_footer=True
 for name,size,bold,fam,col in [('Normal',10.5,False,SERIF,INK),('Title',44,True,SANS,GREEN),('Subtitle',15,False,SANS,GOLD),('Heading 1',16,True,SANS,GREEN),('Heading 2',11.5,True,SANS,GREEN)]:
  set_font(doc.styles[name],size,bold,fam,col)
 for name in ['Heading 1','Heading 2']:
  s=doc.styles[name];s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 8);s.paragraph_format.space_after=Pt(5)
 for st in doc.styles:
  for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 doc.styles['Subtitle'].font.italic=False
 doc.styles['Normal'].paragraph_format.widow_control=True
 hp=sec.header.paragraphs[0];hp.text='旗舰元年  /  完整规则书';run_font(hp.runs[0],SANS,8,GREEN)
 hp.paragraph_format.tab_stops.add_tab_stop(Mm(174),WD_ALIGN_PARAGRAPH.RIGHT)
 r=hp.add_run('\t5.0  综合测试版');run_font(r,SANS,8,GOLD)
 fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run_font(fp.add_run('旗舰元年   ·   '),SANS,8,GREEN)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)
 doc.core_properties.title='旗舰元年 5.0 完整规则书';doc.core_properties.subject='五厂标签、大数值、多源科技与36格双轨；从4.8延续'
 doc.core_properties.author='';doc.core_properties.keywords='桌游,旗舰元年,5.0,测试版'
 settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
 def p(txt='',style=None):return doc.add_paragraph(txt,style)
 def tab(headers,rows,widths=None,fontsize=9.2):
  n=len(headers); width=174
  widths=widths or ([31,143] if n==2 else [38,58,78] if n==3 else [38,48,42,46] if n==4 else [32,31,37,37,37] if n==5 else [44]+[26]*(n-1))
  compact = headers[0]=='定价'
  if compact and doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
  t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  for c,w in zip(t.columns,widths):c.width=Mm(w)
  bp=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','left','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CFDAD2');bp.append(el)
  t._tbl.tblPr.append(bp)
  for ri,vals in enumerate([headers]+rows):
   row=t.rows[0] if ri==0 else t.add_row();rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
   if ri==0:rp.append(OxmlElement('w:tblHeader'))
   for cell,txt,w in zip(row.cells,vals,widths):
    cell.width=Mm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),GREEN if ri==0 else ('F2F5F1' if ri%2==0 else 'FFFFFF'));cp.append(sh)
    mar=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     a=OxmlElement('w:'+edge);a.set(qn('w:w'),'85');a.set(qn('w:type'),'dxa');mar.append(a)
    cp.append(mar);q=cell.paragraphs[0];q.paragraph_format.line_spacing=1.1;q.paragraph_format.space_before=Pt(1.5);q.paragraph_format.space_after=Pt(1.5)
    if ri==0 or (compact and ri<len(rows)):q.paragraph_format.keep_with_next=True
    run_font(q.add_run(str(txt)),SANS if ri==0 else SERIF,fontsize,'FFFFFF' if ri==0 else INK,ri==0)
  q=p();q.paragraph_format.space_after=Pt(2);q.paragraph_format.line_spacing=Pt(2);q.paragraph_format.space_before=Pt(0)
  return t
 # Cover: typography first, no random artwork.
 q=p('FLAGSHIP ERA   /   5.0');q.paragraph_format.space_before=Pt(34);run_font(q.runs[0],SANS,12,GOLD,True)
 q=p('旗舰元年','Title');q.paragraph_format.space_before=Pt(32);q.paragraph_format.space_after=Pt(12)
 p('完整规则书','Subtitle')
 q=p('五厂协同  ·  供应链竞争  ·  技术成果');q.paragraph_format.space_before=Pt(18);run_font(q.runs[0],SANS,13,GREEN)
 q=p('2—4人     90—150分钟（待实测）');q.paragraph_format.space_before=Pt(28);run_font(q.runs[0],SANS,10,GOLD)
 q=p('公开抢到的零件，只是别人看得见的一半。\n真正的旗舰，在共同发布时才揭晓。');q.paragraph_format.space_before=Pt(20);q.paragraph_format.space_after=Pt(22);run_font(q.runs[0],SERIF,16,GREEN)
 tab(['保留骨架','本次整合','配套图鉴'],[['五行动强度\n双轨与共同发布','五厂标签＋多源科技\n从4.8继续迭代','302种内容牌\n按份数准备374张']], [58,58,58],10)
 q=p('综合测试版  /  2026年9月20日');q.paragraph_format.space_before=Pt(20);run_font(q.runs[0],SANS,10,GOLD)
 q=p('科技0—36；前八步每步跨2格，其后每步跨3格。\n以双轨相遇判终局；不使用纸笔。尚待多人平衡实测。');run_font(q.runs[0],SANS,9,'65766A')
 doc.add_page_break()
 # Contents; field references render actual page numbers in LO/Word.
 p('阅读导航','Heading 1')
 p('初次游玩先读第一至第十三节；遇到结算疑问查第十四节。全卡数值与完整效果见配套图鉴，第十六节提供编号和供货批次索引。')
 headings=re.findall(r'^## (.+)$',rules,re.M)
 for idx,h in enumerate(headings,1):
  q=p();q.paragraph_format.space_after=Pt(6);q.paragraph_format.tab_stops.add_tab_stop(Mm(169),WD_ALIGN_PARAGRAPH.RIGHT)
  hyperlink(q,h,anchor=f'ch_{idx}')
  q.add_run('\t');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),f'PAGEREF ch_{idx} \\h');q._p.append(f)
 p('四条重要界线','Heading 2')
 p('类别与厂家分别统计；配件不提供公司厂标；供应世代不等于个人科技；现金不等于收入。')
 p('更新说明','Heading 2')
 p('旧版未改动规则已融入正文，不需要边玩边翻旧版补丁。所有数值为测试参数，未引入电子游玩功能。')
 doc.add_page_break()
 # Intro from source MD excluding repeated title/subtitles.
 pre=rules.split('## 一 ')[0].splitlines()
 for ln in pre[4:]:
  if ln.strip():p(ln)
 lines=rules[rules.index('## 一 '):].splitlines();i=0;chapter=0
 while i<len(lines):
  ln=lines[i].strip()
  if not ln:i+=1;continue
  if ln.startswith('## '):
   chapter+=1;q=p(ln[3:],'Heading 1')
   a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(chapter));a.set(qn('w:name'),f'ch_{chapter}');q._p.insert(0,a)
   z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(chapter));q._p.append(z)
  elif ln.startswith('### '):p(ln[4:],'Heading 2')
  elif ln=='![五厂标志](../assets/factory_strip.png)':
   t=doc.add_table(rows=1,cols=5);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
   for cell,factory in zip(t.rows[0].cells,d['factories']):
    cell.width=Mm(34.8);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F4F5EC');cell._tc.get_or_add_tcPr().append(sh)
    q=cell.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_before=Pt(8);q.paragraph_format.space_after=Pt(4);q.add_run().add_picture(str(A/f"factory_{factory['id']}.png"),width=Mm(11))
    q=cell.add_paragraph(factory['name']+' · '+factory['symbol']);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(7);run_font(q.runs[0],SANS,8,GREEN)
   p()
  elif ln.startswith('https://'):
   q=p();hyperlink(q,'打开出版社官方参考文件',url=ln)
  elif ln.startswith('|'):
   arr=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',v) for v in cells):arr.append(cells)
    i+=1
   n=len(arr[0]);widths=None
   if n==4 and '路线' in arr[0][0]:widths=[32,54,44,44]
   if n==4 and '世代' in arr[0][0]:widths=[22,92,30,30]
   if n==5:widths=[25,37,37,37,38]
   if n==2 and arr[0][0]=='类别':widths=[30,144]
   tab(arr[0],arr[1:],widths);continue
  else:p(ln)
  i+=1
 doc.save(O/'旗舰元年_5.0_完整规则书.docx')

CAT_CSS=(A/'catalog.css').read_text()+FACTORY_CSS+'''
.tagsline{display:flex;gap:5px;align-items:center;flex-wrap:wrap;margin-bottom:2mm;min-height:6mm}.tagsline .neutral{font-size:10px;color:#7c877c}.card.specialProject{--accent:#9a6336}.card h3{margin-bottom:1.6mm}.card .gate{line-height:1.5;margin-bottom:1.6mm}.card .effect{font-size:11.5px;line-height:1.6}.card .effect.dense{font-size:11px;line-height:1.55}.card .life{padding-top:1.3mm;font-size:8.5px}.card.tech .metrics{grid-template-columns:repeat(4,1fr)}.tier b{font-size:17px}.tier{margin:0;padding:2mm .5mm}.tiergrid{margin:2mm 0}.factory-matrix{font-size:10px!important}.factory-matrix td{padding:1.5mm!important}.tight p{font-size:10.5px;line-height:1.85}.align-grid{display:grid;grid-template-columns:repeat(4,1fr);gap:1.6mm;margin:5mm 0}.align-grid>div{background:#edf3ed;border:1px solid #cedbcf;padding:1.6mm;font-size:11px;text-align:center}.newflag{font-size:8px;color:#a17a43}.print-note{font-size:10px;color:#66745f;line-height:1.9}section.legend .counts{margin-top:5mm}
'''

def tags(arr):return ''.join(f'<span class="tag">{E(t)}</span>' for t in arr)
def metrics(vals):return '<div class="metrics">'+''.join(f'<div class="metric">{E(k)}<b>{v}</b></div>' for k,v in vals)+'</div>'

def card(c,typ):
 cid=c['id'];q=c.get('qty',1)
 names={'phone':'产品企划','tech':'科技卡','business':'企业卡','common':'稳定配件','custom':'定制配件','starter':'初始稳定','project':'基础产业突破','specialProject':'专项突破','goal':'终局目标'}
 badge_text=f'{ROM[c["era"]]}代 · {q}张' if typ in ['common','custom'] else f'{names[typ]} · {q}张'
 s=f'<article class="card {typ}" id="card-{cid}" data-card-id="{cid}" data-card-type="{typ}"><div class="card-top"><b>{cid}</b><span>{badge_text}</span></div><h3 class="'+('long' if len(c['name'])>13 else '')+'">'+E(c['name'])+'</h3>'
 if typ in ['phone','tech','business']:
  s+='<div class="tagsline">'+(''.join(badge(f) for f in c.get('factoryTags',[])) or '<span class="neutral">无厂标</span>')+tags(c['tags'])+'</div>'
  if typ=='phone':
   s+=f'<div class="subline">市场：{c["market"]}</div>'
   s+=metrics([('组装',c['strength']),('科技',c['tech']),('开发',c['development']),('验收',c['appeal']),('收入',c['income'])])
   s+=f'<div class="gate">发布前公司门槛　类别：{category_req(c)}；厂家：{factory_req(c)}</div>'
   life='发布开始前查入场门槛；验收并付款成功后取得收入与效果。上市保留，旧机不重复结算。'
  elif typ=='tech':
   s+=metrics([('行动',c['strength']),('科技门槛',c['tech']),('研发费',c['cost']),('获得科技',c['gain'])])
   n=c.get('priorTechCardsRequired',0)
   s+=f'<div class="gate">{c["researchClass"]}　厂家前置：{factory_req(c)}<br>研发基础：'+(f'此前已打出至少{n}张科技卡' if n else '无额外科技卡张数前置')+'</div>'
   life='先检查前置，再付款得印刷科技。自己不计入自身前置；能力与公司标签持续保留。'
  else:
   s+=metrics([('行动',c['strength']),('科技门槛',c['tech']),('合作费',c['cost'])])
   s+=f'<div class="gate">打出前厂家门槛：{factory_req(c)}</div>'
   life='通过合作打出并持续保留。每场发布型效果每张每场至多一次；最终费用最低3现金。'
  s+=f'<p class="effect '+('dense' if len(c['effect'])>88 else '')+'">'+E(c['effect'])+'</p>'
 elif typ in ['common','custom','starter']:
  cs=c['components'];k=KINDS[cs[0]['kind']] if len(cs)==1 else f'{len(cs)}项整合配件'
  s+=f'<div class="subline">{k}　'+('开局分发' if typ=='starter' else '公开供应 · 永久入库' if typ=='common' else '秘密手牌 · 一次性')+'</div>'
  s+='<table class="parttable"><colgroup><col style="width:22%"><col style="width:13%"><col style="width:13%"><col style="width:16%"><col style="width:13%"><col style="width:23%"></colgroup><thead><tr><th>部件</th><th>制造</th><th>吸引力</th><th>功／供</th><th>科技</th><th>类别</th></tr></thead><tbody>'
  for v in cs:
   form='折叠' if v['form']=='fold' else ('直板' if v['kind'] in ['screen','body'] else '')
   label=KINDS[v['kind']]+('·'+form if form else '')
   s+=f'<tr><td>{label}</td><td><b>{v["cost"]}</b></td><td><b>{v["appeal"]}</b></td><td>'+('供' if v['kind']=='body' else '耗')+f'{v["load"]}</td><td>{v["tech"]}</td><td>{v["tag"]}</td></tr>'
  s+='</tbody></table>'
  name=f'<strong>{E(c["abilityName"])}：</strong>' if c.get('abilityName') else ''
  s+='<p class="effect '+('dense' if len(cs)>=3 or len(c['effect'])>78 else '')+'">'+name+E(c['effect'])+'</p>'
  life='每次发布均付制造费；实体卡本场仅用于一机，发布后归库。无公司厂标。' if typ!='custom' else '整张全部部件同机，不拆分、不重叠；逐项付费/查科技。本场后消耗，失败也移出。无厂标。'
 elif typ in ['project','specialProject']:
  s+=f'<div class="subline">统计对象：{E(c["label"])}'+('数量' if c['metric']=='tag' else '')+'</div>'
  s+='<div class="tiergrid">'+''.join(f'<div class="tier"><small>条件至少</small><b>{n}</b><span>取得{v}科技</span></div>' for n,v in zip(c['thresholds'],c['rewards']))+'</div>'
  s+='<p class="effect dense">'+E(c['effect'])+'</p>'
  life=('进入主牌；未提出占手牌。提出者先付6；不提供厂标。每人每项目仅一次，完成推进发布2。' if typ=='specialProject' else '开局随机用4张，其余移出本局。档位全桌各一人，不消耗条件；每人每项目仅一次。')
 else:
  s+='<div class="goalmark">6 <small>分上限 · 仅终局计分</small></div><p class="effect">'+E(c['effect'])+'</p>'
  life='开局抽2留1，秘密保存。不提供类别/厂标，不参与提前触发终局。'
 return s+'<div class="life">'+life+'</div></article>'

def catalog():
 groups=[('phone','产品企划',[c for c in d['cards'] if c['type']=='phone']),('tech','科技卡',[c for c in d['cards'] if c['type']=='tech']),('business','企业卡',[c for c in d['cards'] if c['type']=='business']),('common','稳定配件',[c for c in d['parts'] if c['type']=='common']),('custom','定制配件',[c for c in d['parts'] if c['type']=='custom']),('starter','初始稳定配件',d['starters']),('project','基础产业突破',d['projects']),('specialProject','专项突破',d['specialProjects']),('goal','终局目标',d['goals'])]
 for _,_,cs in groups:cs.sort(key=lambda c:(c['id'][0],int(c['id'][1:])))
 contentpages=sum(math.ceil(len(cs)/6) for _,_,cs in groups)
 total=contentpages+7 # cover/index + 2 action + pricing/factory/track&reference
 index=[];start=3
 for typ,title,cs in groups:
  n=math.ceil(len(cs)/6);index.append((typ,title,len(cs),sum(c.get('qty',1) for c in cs),start,start+n-1));start+=n
 sheets=[];manifest=[]
 def sheet(body,title='',cls='cards-page',anchor=''):
  pn=len(sheets)+1
  hd='' if cls=='cover' else f'<div class="page-head"><h2>{title}</h2><span>旗舰元年 / 5.0</span></div>'
  foot=f'<div class="page-foot"><span>5.0 五厂协同测试版 · 完整数值以本版为准</span><span>{pn} / {total}</span></div>'
  sheets.append(f'<section class="sheet {cls}" id="{anchor or "page-"+str(pn)}">{hd}{body}{foot}</section>')
 cover='<div class="eyebrow">FLAGSHIP ERA / CARD COMPENDIUM 5.0</div><h1>旗舰元年</h1><div class="subtitle">全卡牌图鉴</div>'+factory_strip_html()+'<div class="tagline">五个厂系，多种装配路线。<br>能造什么，与怎样经营，是两种不同的积累。</div><div class="counts"><div><b>302</b>内容牌设计</div><div><b>374</b>内容实体牌</div><div><b>36</b>科技轨完整格数</div></div><div class="panel">企划、科技、企业、稳定、定制、初始配件、基础项目、专项与目标，逐种完整列出。相同编号按份数准备。另附五种行动卡的普通/升级面。</div><small>同一玩家的收入与科技标记相遇或越过触发最后一轮。<br>科技0—8每步跨2收入格；8—36每步跨3格。<br>本图鉴为卡面式参考排版，不是含统一卡背、出血的量产印刷稿。<br>实际玩法见配套完整规则书；尚待多人平衡测试。</small>'
 sheet(cover,cls='cover',anchor='catalog-cover')
 idx='<p>厂标与类别分开显示；“提供厂标”在卡名下，“需求”在门槛栏。需求里的厂名不算图标产出。起始P01—P04没有厂家，通用牌没有厂家并非通配。</p><table><tr><th>类别</th><th>设计／实体</th><th>页码</th></tr>'
 for typ,title,n,q,a,b in index:idx+=f'<tr><td><a href="#cat-{typ}">{title}</a></td><td>{n}／{q}</td><td>{a}—{b}</td></tr>'
 idx+='</table><h3>读数约定</h3><p>企划“验收”是要求四项印刷吸引力之和，不是额外吸引力。“科技门槛”是已有科技要求，“获得科技”才是打出时的收益。配件“制造”每次发布都要付；机身列“供”是供电，其余“耗”是功耗。</p><p>每个企划类别匹配＋2；厂标不自动加吸引力。有效验收最多降低4；打科技、企业与整机各最低付3现金。四类部件必须恰好齐全，折叠屏与折叠机身必须配套。</p><h3>新牌与完整底稿</h3><p>新增T37—T41、B37—B41、R01—R06、O13—O18，共22种。其他原编号和份数延续4.8，但数值已重标，不能混用旧卡。</p><p>开局主牌204张；各代定制陆续加入后，非起始主牌累计244张。稳定80张独立供货；J项目与O目标不混主牌，R专项混主牌。</p>'
 sheet(idx,'图鉴索引与读数',cls='legend',anchor='catalog-index')
 for typ,title,cs in groups:
  for off in range(0,len(cs),6):
   pn=len(sheets)+1
   for c in cs[off:off+6]:manifest.append({'id':c['id'],'type':typ,'name':c['name'],'qty':c.get('qty',1),'page':pn})
   sheet('<div class="grid">'+''.join(card(c,typ) for c in cs[off:off+6])+'</div>',title,anchor='cat-'+typ if off==0 else '')
 acts=[]
 for a in d['actions']:
  for side,label in [('front','普通面'),('back','升级面')]:
   acts.append(f'<article class="card action" data-action-id="{a["id"]}-{side}"><div class="card-top"><b>{a["id"]}</b><span>{label} · 每人1张双面卡</span></div><h3>{a["name"]}</h3><div class="gate">强度S＝当前所在位置（1—5）</div><p class="effect">{E(a[side])}</p><div class="life">用后移回1，其左侧行动顺移。可改整备：仅得9现金，不触发原行动能力。卡面不取代完整边界裁定。</div></article>')
 for off in range(0,10,6):sheet('<div class="grid">'+''.join(acts[off:off+6])+'</div>','五种双面行动卡',anchor='cat-action' if off==0 else '')
 aid='<h3>代表机的市场收入与销售现金</h3><p>格内“收入／现金”；每公司每市场仅一次，不另外追加旧3／1奖励。非代表机只领企划印刷收入与明确自身效果。</p><table><tr><th>定价</th><th>吸引力</th><th>第一</th><th>第二</th><th>第三／第四</th></tr>'
 for p in d['pricing']['strategies']:aid+=f'<tr><td>{p["name"]}</td><td>{p["appealDelta"]:+d}</td>'+''.join(f'<td>{p["incomeByRank"][i]}／{p["cashByRank"][i]}</td>' for i in range(3))+'</tr>'
 aid+='</table><h3>扣牌锁定，无需纸笔</h3><p>每人9张定价牌，三个筹备位各一套三张。发布前同时秘密选牌，与配置一起翻开；不额外行动，不允许改配置。漏放、多放或编号不符按标准价。结算后全部收回。</p><p>调校标记放在本机实际提供能力的配件上；T19放对应筹备位的“T19节能”格。不放为默认，每机最多一个方案，只有正常返工可重选。</p><h3>现金与科技速查</h3><table><tr><th>项目</th><th>5.0数值</th></tr><tr><td>初始现金／收入／科技</td><td>60／5／0</td></tr><tr><td>正常发布回款</td><td>20＋当前收入</td></tr><tr><td>基础研究</td><td>普通S4付24；升级S3付18，得1科技</td></tr><tr><td>基础研究次数</td><td>每研发行动一次，整局可重复</td></tr><tr><td>供应II／III／IV／V</td><td>科技2／4／6／8</td></tr><tr><td>升级／第三筹备位</td><td>科技3、7升级；科技5且上市2张开第三位</td></tr></table><p>最终发布仍结算价格收入与现金，但不发正常回款。所有制造费须在任何本场奖励之前支付。</p>'
 sheet(aid,'定价、现金与科技速查',cls='aid',anchor='cat-aid')
 factory=factory_strip_html()+'<h3>公司积累与产品积累不是同一对象</h3><p>公司厂标：已上市企划＋已打科技＋已打企业。产品厂系：只看已上市企划。项目、目标、配件不提供厂标；无厂标不是通配。</p><h3>五厂分布</h3><table class="factory-matrix"><tr><th>厂系</th><th>企划</th><th>科技</th><th>企业</th><th>起始固定阵营</th></tr>'
 for f in d['factories']:
  vals=[sum(f['id'] in c.get('factoryTags',[]) for c in d['cards'] if c['type']==typ) for typ in ['phone','tech','business']]
  factory+=f'<tr><td>{badge(f["id"])}</td>'+''.join(f'<td>{v}</td>' for v in vals)+'<td>无</td></tr>'
 factory+='</table><h3>专项提出与竞争</h3><p>用合格合作行动，从手牌提出1张R。必须立刻满足一个档位，先付6现金，再公开并完成。之后留在公共区供其他玩家竞争，不是提出者的公司牌。每玩家每项目整局一次，不能补差升级；完成均推进发布2。</p><h3>多源科技</h3><p>13张应用技术不直接给点，28张成果技术给1—3科技，全体科技卡印刷合计36点。27张企划有条件科技奖励。J基础项目、R专项、可重复基础研究以及T39/T41新成果引擎提供其他路径。</p><p>有来源不等于每局都能全部取得。T39/T41必须实际发布合格新机，空发布和旧机不产生科技；不能把所有厂标混进类别种类。</p>'
 sheet(factory,'五厂标签与研发路径',cls='legend',anchor='cat-factory')
 align='<p>每项表示“科技格 → 收入对齐线”。收入到该线即相遇，越过即触发；完整结束步骤见规则书。版图已按此定位，不要求边玩边换算。</p><div class="align-grid">'+''.join(f'<div>科技 {t} → 收入 {v}</div>' for t,v in enumerate(d['scoreTrack']['alignment']))+'</div><h3>终局读分</h3><p>科技对齐线是零分线。收入超过几格就几分，未到差几格就负几分，再加目标。科技超过36的部分每点另计3基础分，不套用余数0—8的短跨度。</p><h3>出处与创作边界</h3><p>规则底稿：本项目4.8完整数据与正文。五厂、专项项目、重标数值和科技来源为本次5.0设计。参照《方舟动物园》出版社规则与图鉴的图标积累、手牌项目、受限投入和成果引擎结构，不照搬动物或品牌。</p><p>芯片功能资料沿用4.7/4.8存档，本轮未再次核验真实采购价、发布日期与性能。厂标不指认为现实品牌；费用、功耗、吸引力全部为桌游单位。</p><p>图鉴已完整列出302种内容牌。程序与排版核对不是多人平衡验证，具体检查范围见随包说明。</p>'
 sheet(align,'双轨对齐与资料边界',cls='legend tight',anchor='cat-track')
 assert len(sheets)==total,(len(sheets),total)
 links=''.join(f'<a href="#cat-{typ}">{title}</a>' for typ,title,_,_,_,_ in index)
 h=f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>旗舰元年 5.0 · 全卡牌图鉴</title><style>{CAT_CSS}</style></head><body><header class="tools"><b>旗舰元年 5.0</b><div class="links">{links}</div><button onclick="window.print()">打印图鉴</button></header>'+''.join(sheets)+'</body></html>'
 (O/'旗舰元年_5.0_全卡牌图鉴.html').write_text(h,encoding='utf8')
 (W/'catalog_manifest.json').write_text(json.dumps({'pages':total,'cards':manifest,'index':index},ensure_ascii=False,indent=2),encoding='utf8')
 md='# 旗舰元年5.0全卡牌图鉴\n\n302种内容牌，374张实体牌；另外准备行动与定价等功能组件。\n\n'
 for typ,title,cs in groups:
  md+='## '+title+'\n\n'
  for c in cs:
   md+=f'### {c["id"]} {c["name"]}（{c.get("qty",1)}张）\n\n'
   if typ in ['phone','tech','business']:
    md+='提供类别：'+'、'.join(c['tags'])+'；提供厂标：'+factory_text(c['factoryTags'])+'。\n\n'
    md+=f'行动{c["strength"]}；科技门槛{c["tech"]}；'
    if typ=='phone':md+=f'市场{c["market"]}；开发费{c["development"]}；验收{c["appeal"]}；收入{c["income"]}；类别门槛{category_req(c)}；厂家门槛{factory_req(c)}。\n\n'
    else:
     md+=f'费用{c["cost"]}；厂家门槛{factory_req(c)}；'
     if typ=='tech':md+=f'{c["researchClass"]}；获得科技{c["gain"]}；此前科技卡要求{c["priorTechCardsRequired"]}张。'
     md+='\n\n'
   elif typ in ['common','custom','starter']:
    md+=f'供应世代：{ROM[c.get("era",0)]}；不提供公司厂标。\n\n'
    for p in c['components']:md+=f'{KINDS[p["kind"]]}：{p["name"]}；制造{p["cost"]}；吸引力{p["appeal"]}；'+('供电' if p['kind']=='body' else '功耗')+f'{p["load"]}；科技{p["tech"]}；类别{p["tag"]}；形态{p["form"]}。\n\n'
   elif typ in ['project','specialProject']:md+=f'统计{c["label"]}；门槛'+ '／'.join(map(str,c['thresholds']))+'；科技奖励'+'／'.join(map(str,c['rewards']))+'；不提供标签。\n\n'
   md+='完整效果：'+c['effect']+'\n\n'
   if typ in ['common','starter']:md+='稳定生命周期：每场每张只供一机；每次仍付制造费；发布后归库，成功失败均同。\n\n'
   elif typ=='custom':md+='定制生命周期：整张同机，不拆不重叠；本场后移出，失败也消耗；发布前返工取回不消耗。\n\n'
 md+='## 五行动两面\n\n'
 for c in d['actions']:md+=f'### {c["id"]} {c["name"]}\n\n普通：{c["front"]}\n\n升级：{c["back"]}\n\n'
 (W/'v50_catalog.md').write_text(md,encoding='utf8')
 print('Catalog',len(manifest),'designs',sum(c['qty'] for c in manifest),'cards',total,'pages')

if __name__=='__main__':
 create_icons();rulebook();catalog()
 print('Generated complete DOCX and HTML catalogue.')
