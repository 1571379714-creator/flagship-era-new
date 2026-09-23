from pathlib import Path
import json,re,html
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[1];O=R/'outputs';O.mkdir(exist_ok=True)
book=json.loads((R/'work/v69_rulebook.json').read_text(encoding='utf8'))
E=lambda x:html.escape(str(x))
def xt(tag,attrs=None):
 x=OxmlElement('w:'+tag)
 for k,v in (attrs or {}).items():x.set(qn('w:'+k),str(v))
 return x

def rich(p,text):
 for s in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
  r=p.add_run(s.strip('*`') if s.startswith(('**','`')) else s)
  if s.startswith('**'):r.bold=True
def paragraphs(md):
 ls=md.strip().splitlines();out=[];i=0
 while i<len(ls):
  if not ls[i].strip():i+=1;continue
  if ls[i].startswith('|'):
   rows=[]
   while i<len(ls) and ls[i].startswith('|'):
    vals=[v.strip() for v in ls[i].strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in vals):rows.append(vals)
    i+=1
   out.append(('table',rows));continue
  out.append(('p',ls[i]));i+=1
 return out
def hrich(t):
 t=E(t);t=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',t);return t
def hmd(md):
 out=[]
 for typ,x in paragraphs(md):
  if typ=='p':out.append('<p>'+hrich(x)+'</p>')
  else:out.append('<table>'+''.join('<tr>'+''.join(('<th>' if i==0 else '<td>')+hrich(v)+('</th>' if i==0 else '</td>') for v in row)+'</tr>' for i,row in enumerate(x))+'</table>')
 return ''.join(out)
def set_cell(cell,width,fill=None):
 cell.width=Mm(width);pr=cell._tc.get_or_add_tcPr()
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
 mar=xt('tcMar')
 for edge,val in [('top',70),('bottom',70),('left',80),('right',80)]:mar.append(xt(edge,{'w':val,'type':'dxa'}))
 pr.append(mar)
 if fill:pr.append(xt('shd',{'fill':fill}))
def add_body(parent,md,width=123):
 for typ,x in paragraphs(md):
  if typ=='p':
   p=parent.add_paragraph();rich(p,x)
   p.paragraph_format.keep_with_next=False
  else:
   n=len(x[0]);t=parent.add_table(rows=1,cols=n);t.autofit=False
   widths=([width*.29,width*.71] if n==2 else [width*.30,width*.17,width*.53] if n==3 else [width/n]*n)
   if n==2 and x[0][0]=='你现在遇到的问题':widths=[width*.76,width*.24]
   if n==3 and x[0][0] in ('参考机定位','参考格竞争力'):widths=[width*.27,width*.24,width*.49]
   grid=t._tbl.tblGrid
   for ch in list(grid):grid.remove(ch)
   for w in widths:grid.append(xt('gridCol',{'w':round(w/25.4*1440)}))
   for i,row in enumerate(x):
    rw=t.rows[0] if i==0 else t.add_row();rw._tr.get_or_add_trPr().append(xt('cantSplit'))
    if i==0:rw._tr.get_or_add_trPr().append(xt('tblHeader'))
    for j,v in enumerate(row):
     cell=rw.cells[j];set_cell(cell,widths[j],'E1EFF0' if i==0 else ('F3F7F8' if i%2==0 else 'FFFFFF'))
     p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.1;rich(p,v)
     for run in p.runs:run.font.size=Pt(9.3);run.bold=i==0
def make_rulebook():
 doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.left_margin=sec.right_margin=Mm(15);sec.top_margin=Mm(17);sec.bottom_margin=Mm(17);sec.header_distance=Mm(7);sec.footer_distance=Mm(8)
 for name,sz,bold in [('Normal',10.5,False),('Title',33,True),('Subtitle',15,False),('Heading 1',17,True),('Heading 2',12.5,True)]:
  st=doc.styles[name];st.font.name='Noto Sans CJK SC' if bold else 'Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),st.font.name);st.font.size=Pt(sz);st.font.bold=bold
  st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.2;st.paragraph_format.widow_control=True
  if 'Heading' in name:st.font.color.rgb=RGBColor.from_string('146A76');st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10)
 p=sec.header.paragraphs[0];p.text='旗舰元年  /  完整规则书                                         6.9  五十平台与芯片专长'
 for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('57707A')
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('FLAGSHIP ERA  /  6.9   ·   ').font.size=Pt(8);p._p.append(xt('fldSimple',{'instr':'PAGE'}))
 doc.add_paragraph('FLAGSHIP ERA  /  6.9',style='Subtitle');doc.add_paragraph('旗舰元年',style='Title');doc.add_paragraph('完整规则书',style='Title');doc.add_paragraph('五十平台与芯片专长',style='Subtitle')
 doc.add_paragraph('\n')
 p=doc.add_paragraph('产品先备好，价格先锁定。\n竞品，发布时才揭晓。');p.runs[0].font.size=Pt(20)
 doc.add_paragraph('\n')
 for a,b in [('经营方式','统一拿牌、四栏供货、一次造一批；卖掉或取消才结束占用。'),('发布时间','由玩家共同推进，没有固定轮数或固定销售季。'),('读法','正文讲怎么做；右栏放例子、实体摆法与关联章节。')]:
  p=doc.add_paragraph();p.add_run(a+'  ').bold=True;p.add_run(b)
 doc.add_paragraph('\n2—4人  ·  315种设计 / 315张内容牌\n2026年9月23日  ·  测试参数，尚未完成充分多人实测。')
 doc.add_page_break()
 doc.add_paragraph('阅读导航',style='Heading 1')
 doc.add_paragraph('先看§1—3认识资源和回合。第一次造机查§4—8；科研授权查§9—10；合同查§11；发布查§12—14。卡牌的完整数值另见图鉴。')
 toc=doc.add_table(rows=1,cols=2);toc.autofit=False
 grid=toc._tbl.tblGrid
 for ch in list(grid):grid.remove(ch)
 for w in [15,165]:grid.append(xt('gridCol',{'w':round(w/25.4*1440)}))
 for i,s in enumerate(book['sections']):
  rw=toc.rows[0] if i==0 else toc.add_row()
  set_cell(rw.cells[0],15);set_cell(rw.cells[1],165)
  rw.cells[0].paragraphs[0].text=s['number'];rw.cells[1].paragraphs[0].text=s['title']
  for ce in rw.cells:
   for p in ce.paragraphs:p.paragraph_format.space_after=Pt(2)
 doc.add_paragraph('正文与侧栏的分工',style='Heading 2');doc.add_paragraph('必须执行的条件、费用和时点都在左侧正文。右侧的例子只帮助理解，不新增隐藏规则。章节编号在Word、PDF和HTML中一致。')
 doc.add_paragraph('四个容易混淆的词',style='Heading 2');doc.add_paragraph('产品位＝一批现货；稳定库＝共享供货。本人回合＝这次3个工作点；准备期＝两次实际发布之间。')
 doc.add_page_break()
 for s in book['sections']:
  doc.add_paragraph(s['number']+'  '+s['title'],style='Heading 1')
  if s['intro']:
   intro=doc.add_paragraph(s['intro']);intro.paragraph_format.keep_with_next=True
  for b in s['blocks']:
   t=doc.add_table(rows=1,cols=3);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
   # 一个短主题完整保留；章节按内容自然流动，不强制一章一页。
   t.rows[0]._tr.get_or_add_trPr().append(xt('cantSplit'))
   widths=[126,5,49]
   grid=t._tbl.tblGrid
   for ch in list(grid):grid.remove(ch)
   for w in widths:grid.append(xt('gridCol',{'w':round(w/25.4*1440)}))
   left,gap,side=t.rows[0].cells
   set_cell(left,126);set_cell(gap,5);set_cell(side,49,'EEF5F6')
   p=left.paragraphs[0];p.style=doc.styles['Heading 2'];rich(p,b['title']);p.paragraph_format.space_before=Pt(0)
   add_body(left,b['body'],122)
   p=side.paragraphs[0];p.add_run(b['asideKind']).bold=True;p.paragraph_format.space_after=Pt(7)
   for line in b['aside'].split('\n'):
    p=side.add_paragraph();rich(p,line);p.paragraph_format.line_spacing=1.3;p.paragraph_format.space_after=Pt(7)
    for r in p.runs:r.font.size=Pt(9.2);r.font.color.rgb=RGBColor.from_string('3C606B')
   for cell in (left,gap,side):
    for p in cell.paragraphs:p.paragraph_format.widow_control=True
   p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run(' ').font.size=Pt(2)
 doc.core_properties.title='旗舰元年6.9 完整规则书';doc.core_properties.author='';doc.core_properties.subject='正文与侧栏 / 隐藏参考 / 实体操作'
 doc.save(O/'旗舰元年_6.9_完整规则书.docx')
 css='''*{box-sizing:border-box}body{font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;color:#223B45;background:#E9F0F3;line-height:1.8;margin:0}header{padding:30px;background:#193E4C;color:white}header h1{margin:0}nav{display:flex;flex-wrap:wrap;gap:10px;padding:20px}nav a{color:#176B79;text-decoration:none}main{max-width:1120px;margin:auto;background:white;padding:32px}h2{border-bottom:2px solid #176B79;color:#176B79}h3{margin-top:0;color:#176B79;font-size:20px}.ruleblock{display:grid;grid-template-columns:minmax(0,2.65fr) minmax(200px,1fr);gap:24px;margin:22px 0 30px}.ruleblock aside{background:#EEF5F6;border-left:3px solid #88ADB7;padding:18px;color:#426371;font-size:14px;align-self:start}.ruleblock p{margin:0 0 12px}table{width:100%;border-collapse:collapse;font-size:15px;margin:12px 0}th,td{border-bottom:1px solid #CCDDE1;padding:8px 9px;text-align:left}th{background:#E4F0F2}button{padding:8px 20px}@media(max-width:750px){main{padding:18px}.ruleblock{display:block}.ruleblock aside{margin-top:12px}table{font-size:13px}}@media print{body{background:white}nav,button{display:none}main{padding:0;max-width:none}.ruleblock{break-inside:avoid}h2,h3{break-after:avoid}header{background:white;color:#176B79}}'''
 nav=''.join(f'<a href="#s{s["number"]}">{s["number"]} {E(s["title"])}</a>' for s in book['sections'])
 body=''
 for s in book['sections']:
  body+=f'<section id="s{s["number"]}"><h2>{s["number"]} {E(s["title"])}</h2><p>{E(s["intro"])}</p>'
  for b in s['blocks']:body+='<div class="ruleblock"><div><h3>'+E(b['title'])+'</h3>'+hmd(b['body'])+'</div><aside><strong>'+E(b['asideKind'])+'</strong>'+''.join('<p>'+E(x)+'</p>' for x in b['aside'].split('\n'))+'</aside></div>'
  body+='</section>'
 (O/'旗舰元年_6.9_完整规则书.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>旗舰元年6.9 易读规则</title><style>'+css+'</style><header><h1>旗舰元年 6.9</h1><p>完整规则书 · 五十平台与芯片专长</p><button onclick="window.print()">打印</button></header><nav>'+nav+'</nav><main>'+body+'</main></html>',encoding='utf8')
