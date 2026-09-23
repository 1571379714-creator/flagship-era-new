"""只从现行JSON/正文导出完整成品；不运行迁移，不改卡牌数据。"""
from pathlib import Path
import json, html, re, os, shutil, subprocess, tempfile, argparse
from doc_layout import make_rulebook, hmd, Document, Mm, Pt, RGBColor, qn, xt, rich, add_body
R=Path(__file__).resolve().parents[1]; O=R/'outputs'; O.mkdir(exist_ok=True)
d=json.loads((R/'work/v70_data.json').read_text(encoding='utf8')); E=lambda x:html.escape(str(x))
M={x['id']:x['name'] for x in d['markets']};F={x['id']:x for x in d['factories']}
K={'chip':'芯片','screen':'屏幕','camera':'影像','body':'机身'}
N={'supply':'稳定配件','custom':'定制配件','starter':'初始通用配件','technology':'研发项目','contract':'合同','facility':'设施','scheme':'产品方案','benchmark':'参考机','demand':'需求'}
C={'supply':'#377D89','custom':'#9A604B','starter':'#627984','technology':'#236F94','contract':'#23765F','facility':'#986B32','scheme':'#74528E','benchmark':'#AA4D47','demand':'#506A9A'}
CSS=(R/'assets/catalog.css').read_text()+'.card td{white-space:pre-line}.sourcehint{font-size:7.0pt!important;color:#627581;margin:1mm 0!important}.card table{font-size:8.1pt;line-height:1.25}.card th,.card td{padding:1mm .6mm}' +'\n.card .metrics{margin:1.4mm 0}.card .small{font-size:7.3pt}.card .featurehint{font-size:7pt;line-height:1.35;color:#64777f}.card .special{font-size:8pt;background:#F6F0E8;padding:1.6mm}.card .long{font-size:8pt;line-height:1.4}\n'
CSS += (R/'assets/supplier_cards.css').read_text(encoding='utf8')
CSS += '.factoryicon{width:3.1mm;height:3.1mm;display:inline-block;vertical-align:middle;margin-right:1mm}.card .researchmodes{margin:1mm 0;font-size:8pt}.researchmodes th,.researchmodes td{padding:.6mm .7mm}.main-grid .card .researchmodes{font-size:7.0pt;line-height:1.15;margin:.6mm 0}.main-grid .researchmodes th,.main-grid .researchmodes td{padding:.35mm .6mm}.main-grid .card .factoryicon{width:2.3mm;height:2.3mm}'
def factory_icon(fid):
 f=F[fid];svg=(R/f'assets/factory_{fid}.svg').read_text(encoding='utf8')
 return svg.replace('<svg ', '<svg class="factoryicon" aria-label="'+E(f['name'])+'" ',1)

def factory_badge(fid,supply=False):
 f=F[fid]
 return f'<span class="factory" style="--fc:{f["color"]}">{factory_icon(fid)}'+E(f['name'])+'</span>'

def metric(label,val):return f'<span class="metric">{label}<b>{val}</b></span>'
def table(headers,rows):return '<table><tr>'+''.join('<th>'+E(h)+'</th>' for h in headers)+'</tr>'+''.join('<tr>'+''.join('<td>'+E(c)+'</td>' for c in row)+'</tr>' for row in rows)+'</table>'
def card_body(c):
 t=c['type'];s=''
 if t in ('supply','custom','starter'):
  s=''
  if c.get('supplier'):
   s+=factory_badge(c['supplier'],True)
  elif t=='starter':s+='<span class="factory neutral">中立初始 · 开局免费入库</span>'
  s+='<div class="metrics">'+metric('世代','初始' if t=='starter' else c['era'])
  if t=='supply':s+=metric('普通入库',c['installCost'])+metric('本厂伙伴',c['installPartnerCost'])
  if c.get('integratedSlots'):s+=metric('集成占位',c['integratedSlots'])
  s+='</div>'+table(['部件（形态）','制造费','规格','耗 / 供','标识'],[[K[a['kind']]+((' / '+('折叠' if a['form']=='fold' else '直板')) if a['kind'] in ('screen','body') else '')+'\n'+a['name'],a['batchCost'],a['spec'],('供' if a['kind']=='body' else '耗')+str(a['power']),a['mark']] for a in c['components']])
  s+='<p class="sourcehint">'+E(('初始发放 · M0'+c['id'][2:3] if t=='starter' else '年代 '+'—'.join(map(str,c.get('eraYears',[]))))+(' · '+str(c.get('chipYear',''))+' '+c.get('chipTier','') if c.get('chipYear') else ''))+'</p>'
  s+=('<p class="sourcehint">'+E(c['cardEvidenceNote'])+'</p>' if c.get('cardEvidenceNote') else '')
  s+='<p class="effect">'+E(c['effect'])+'</p><p class="small">'+E(c['lifecycle'])+'</p>'
  hints=[]
  kinds={a['kind'] for a in c['components']}
  if 'chip' in kinds:hints.append('初始通用平台；性能≥6，生态读标识。' if t=='starter' else '规格为压缩档，功耗非瓦数；性能≥6，生态读标识。')
  if 'camera' in kinds:hints.append('影像：影像规格≥6，除非能力另授予。')
  if 'body' in kinds:hints.append('续航：余量≥4；轻薄：机身轻薄标识。')
  if 'screen' in kinds or 'body' in kinds:hints.append('折叠屏＋折叠机身必须配套。')
  s+='<p class="featurehint">'+E(' '.join(hints))+'</p>'
 elif t=='technology':
  if c['jointEligible']:
   s='<table class="researchmodes"><tr><th>立项模式</th><th>现金</th><th>工期</th></tr><tr><td>联合</td><td>'+str(c['jointCost'])+'</td><td>'+str(c['jointWork'])+'</td></tr><tr><td>自主</td><td>'+str(c['cost'])+'</td><td>'+str(c['selfWork'])+'</td></tr></table>'
   s+='<div class="metrics">'+metric('完成科技',c['gain'])+metric('授权/批',c['licenseFee'])+'</div>'
  else:
   s='<div class="metrics">'+metric('自主费',c['cost'])+metric('自主工期',c['selfWork'])+metric('科技',c['gain'])+'</div>'
  s+=f'<p class="small">领域：{E(c["field"])} · 本人此前完成≥{c["prerequisiteProjects"]}</p>'
  s+=f'<p class="proof"><b>证据：</b>{E(c["proof"])}</p><p class="effect">{E(c["effect"])}</p>'
  if c['jointEligible']:
   s+=f'<p class="jointnote">联合：限核心厂立项，完成另需含本厂{K[c["relatedPart"]]}的合法样机。使用关联：<b>{K[c["relatedPart"]]}同厂</b>；伙伴可按批付作者，作者自用免费。外部授权不算本人科研。</p>'
   s+='<div class="licensewells"><div class="authorwell">圆形<br>作者</div>'+''.join(f'<div class="userwell"><b>{a}</b><small>产品片</small></div>' for a in 'ABCD')+'</div>'
  else:s+='<p class="jointnote">公司工艺 · 仅自主研发 · 不公开授权 · 不占产品技术位</p>'
 elif t=='contract':
  s=factory_badge(c['factory'])
  s+='<div class="metrics">'+metric('期限',str(c['periods'])+'次发布')+metric('押金' if c.get('deposit') else '费用',c.get('deposit') or c.get('fee',0))+'</div>'
  if c['subtype']=='order':s+='<p><b>整单 '+str(c['units'])+'批：'+str(c['totalCash'])+'现金 / '+str(c['totalIncome'])+'收入；另退押'+str(c['deposit'])+'。</b></p>'
  s+='<p class="small"><b>条件：</b>'+E(c['requirement'])+'</p><p class="effect long">'+E(c['effect'])+'</p>'
 elif t=='facility':
  s='<div class="metrics">'+metric('建设费',c['cost'])+metric('行动','合作')+metric('空间',c['space'])+'</div><p class="effect">'+E(c['effect'])+'</p><p class="small">先从主牌取得，再用合作行动付现金建设。限次片放本卡；取消产品不复位。</p>'
 elif t=='scheme':
  s='<div class="metrics">'+metric('方案费',c['developmentCost'])+metric('每批最多','1方案')+'</div><p><b>要求：</b>'+E(c['requirement'])+'</p><p class="effect">'+E(c['effect'])+'</p><p class="small">'+E(c['lifecycle'])+'</p>'
 elif t=='benchmark':
  s='<div class="metrics">'+metric('参考代',c['era'])+metric('各市场占','1订单')+'</div>'
  rows=[]
  for m in c['markets']:rows.append([M[m['market']],m['score'],m['reward'],('—' if m['reward']==0 else f'{m["incomeReward"]} / {m["technologyReward"]}'),m['specialText']])
  s+=table(['挑战市场','竞争力','星','收入 / 科技','特色'],rows)
  s+='<p class="special"><b>'+E(c['specialName'])+'：</b>'+E(c['specialSummary'])+'</p><p class="small effect">'+E(c['effect'])+'</p>'
 elif t=='demand':
  s=table(['市场','预算','订单2/3/4人','偏好（各＋3）'],[[M[m['market']],m['budget'],' / '.join(str(m['ordersByPlayers'][str(n)]) for n in (2,3,4)),' / '.join(m['preferences'])] for m in c['markets']])
  s+='<p class="effect">'+E(c['effect'])+'</p><p class="small">入场：大众任意合法产品；电竞需性能；摄影需影像；商务需续航/生态/折叠之一。能力可另许可。</p>'
 return s

def render_card(c):
 foot='主牌 · 五代开局全混' if c.get('deck')=='main' else ('初始发放，不混主牌' if c['type']=='starter' else '功能牌，不入手牌')
 if c['type']=='technology':foot+=' · '+('公司工艺' if c['scope']=='process' else '普通每批2技术位')
 return f'<article class="card" id="{c["id"]}" data-type="{c['type']}" style="--accent:{C[c["type"]]}"><div class="eyebrow"><span>{c["id"]} / {N[c["type"]]}</span><span>{str(c["era"])+"代 · " if c.get("deck")=="main" else ""}实体 ×{c["qty"]}</span></div><h3>{E(c["name"])}</h3>{card_body(c)}<div class="bottom"><span>{foot}</span><span>7.0</span></div></article>'

def document_html(title,style,body,header=''):
 return '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+E(title)+'</title><style>'+style+'</style><body><header class="web"><h1>'+E(title)+'</h1><button onclick="window.print()">打印</button>'+header+'</header>'+body+'</body></html>'

def build_catalog():
 groups=[('稳定配件',[c for c in d['parts'] if c['type']=='supply'],4),('定制配件',[c for c in d['parts'] if c['type']=='custom'],4),('初始通用配件',[c for c in d['parts'] if c['type']=='starter'],4),('研发项目',d['technologies'],4),('经营合同',d['contracts'],6),('设施建设',d['facilities'],6),('产品方案',d['schemes'],6),('参考机',d['benchmarks'],4),('需求牌',d['demands'],4)]
 pages=[];manifest=[];ranges=[];page=3
 for title,cards,n in groups:
  start=page
  for i in range(0,len(cards),n):
   cs=cards[i:i+n];pages.append(f'<section class="page"><div class="pagehead"><b>{title}</b><span>7.0 / 完整图鉴</span></div><div class="cardgrid {"tall" if n==4 else ""}">'+''.join(render_card(c) for c in cs)+f'</div><div class="pagefoot"><span>统一牌流 · 单批产品 · 供应商网络 · 技术署名与按批授权 · 六卡双行 · 对向双轨 · 完整7.0卡文</span><span>{page}</span></div></section>')
   for c in cs:manifest.append({'id':c['id'],'name':c['name'],'page':page,'qty':c['qty'],'type':c['type']})
   page+=1
  ranges.append([title,len(cards),start,page-1])
 cover='<section class="page cover"><p>FLAGSHIP ERA / 7.0</p><div class="big">旗舰元年</div><div class="sub">完整卡牌图鉴</div><div class="lede">取得免费，稳定入库付费。<br>四栏供货，代替卡引用；一组配置就是一批现货。<br>读卡上的数值，做桌上的生意。</div><div class="statrow"><div><strong>315</strong>种设计</div><div><strong>315</strong>张内容牌</div><div><strong>257</strong>张主牌</div></div><p class="watermark">完整卡面参考稿 · 每种实体均一张<br>本图鉴按种展示；另附257张统一主牌正反面。<br>不含量产出血刀模；功能组件另册。<br>2026年9月23日 · 未充分多人实测</p></section>'
 index='<section class="page index"><h2>内容索引</h2>'+table(['类别','设计数','页码'],[[x[0],x[1],f'{x[2]}—{x[3]}'] for x in ranges])+'<div class="note">主牌：稳定80、定制64、研发41、合同30、设施18、方案24，共257张。五代开局全部混洗。<br>取牌不另收现金；稳定读普通/伙伴入库价，每批制造另付。<br>研发的工期与科技奖励不是同一数字。合同期限按发布，不按玩家回合。<br>稳定实体不离库；取消回件、成交消耗。扩产同组加一批，实际卡与指派只清理一次。</div><p>需求的订单量直接印2/3/4人列；参考每格直接印收入、科技和特色值。研发自主/联合及授权费直接印卡；作者圆片永久、使用指派随批次结束收回。</p></section>'
 (O/'旗舰元年_7.0_全卡牌图鉴.html').write_text(document_html('旗舰元年 7.0 · 全卡牌图鉴',CSS,cover+index+''.join(pages)),encoding='utf8')
 (R/'work/catalog_manifest.json').write_text(json.dumps({'cards':manifest,'pages':page-1,'ranges':ranges},ensure_ascii=False,indent=2),encoding='utf8')
 text=['# 旗舰元年7.0 完整文字图鉴','315种设计 / 315张内容牌；来源work/v70_data.json。']
 for title,cs,n in groups:
  text+=['\n## '+title]
  for c in cs:
   text+=['\n### '+c['id']+' '+c['name']+'（×'+str(c['qty'])+'）']
   text += [k+'：'+(json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else str(v)) for k,v in c.items() if k not in ('id','name')]
 (R/'work/v70_catalog.md').write_text('\n\n'.join(text)+'\n',encoding='utf8')

def build_planning():
 from docx.enum.text import WD_ALIGN_PARAGRAPH
 doc=Document();s=doc.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(18);s.left_margin=s.right_margin=Mm(18)
 for n,size,b in [('Normal',10.5,False),('Title',29,True),('Heading 1',17,True),('Heading 2',12,True)]:
  st=doc.styles[n];st.font.name='Noto Sans CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),st.font.name);st.font.size=Pt(size);st.font.bold=b;st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.13
  if n.startswith('Heading'):st.font.color.rgb=RGBColor.from_string('146A76');st.paragraph_format.keep_with_next=True
 s.header.paragraphs[0].text='旗舰元年 / 7.0 设计规划书';s.header.paragraphs[0].runs[0].font.size=Pt(8)
 p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('7.0 / ');p._p.append(xt('fldSimple',{'instr':'PAGE'}))
 doc.add_paragraph('7.0 双行排程设计规划',style='Title');doc.add_paragraph('旗舰元年 · 从6.9结构迁移 · 六行动、对向双轨与节奏校准')
 text=(R/'work/v70_planning.md').read_text(encoding='utf8')
 for typ,x in __import__('doc_layout').paragraphs(text):
  if typ=='table':
   # add_body works with cell; place a borderless parent cell for structured table.
   holder=doc.add_table(rows=1,cols=1);holder.autofit=False;holder.cell(0,0).width=Mm(174)
   md='\n'.join('| '+' | '.join(row)+' |' for row in x);add_body(holder.cell(0,0),md,170)
  elif x.startswith('# '):continue
  elif x.startswith('## '):doc.add_paragraph(x[3:],'Heading 1')
  elif x.startswith('### '):doc.add_paragraph(x[4:],'Heading 2')
  else:rich(doc.add_paragraph(),x)
 doc.core_properties.author='';doc.core_properties.title='旗舰元年7.0设计规划书';doc.save(O/'旗舰元年_7.0_设计规划书.docx')

# Native HTML layout boards and reusable functional cards; no game state functions.
from physical_layout import build_reference_print,build_main_print,build_starter_print
from surfaces_v70 import build_boards,build_aids,build_action_cards
from supplier_layout import build_supplier_aids
from expansion_layout import build_expansion_aids

def export_pdf():
 lo=shutil.which('libreoffice') or shutil.which('soffice')
 if not lo:raise RuntimeError('需要LibreOffice导出规则书与规划书PDF')
 for f in O.glob('*.docx'):
  profile=Path(tempfile.mkdtemp(prefix='flagship70-lo-')).as_uri()
  subprocess.run([lo,f'-env:UserInstallation={profile}','--headless','--convert-to','pdf','--outdir',str(O),str(f)],check=True,timeout=120,capture_output=True)
 from playwright.sync_api import sync_playwright
 with sync_playwright() as pw:
  browser=pw.chromium.launch(executable_path=os.environ.get('FLAGSHIP_CHROMIUM') or shutil.which('chromium'),args=['--no-sandbox'])
  for f in O.glob('*.html'):
   if f.with_suffix('.docx').exists():continue
   p=browser.new_page();p.set_content(f.read_text(encoding='utf8'),wait_until='load');p.evaluate('document.fonts.ready')
   if '线下版图' in f.name:p.evaluate('document.body.classList.add("printall")')
   p.pdf(path=str(f.with_suffix('.pdf')),prefer_css_page_size=True,print_background=True);p.close()
  browser.close()
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--no-pdf',action='store_true');args=ap.parse_args()
 from sync_rules_text import check
 check();make_rulebook();build_planning();build_catalog();build_boards(globals());build_aids(globals());build_reference_print(globals());build_main_print(globals());build_starter_print(globals());build_supplier_aids(globals());build_expansion_aids(globals());build_action_cards(globals());__import__('build_chip_report').build_chip_report()
 if not args.no_pdf:export_pdf()
 print('Built 7.0 actual files; no migration or remote write.')
