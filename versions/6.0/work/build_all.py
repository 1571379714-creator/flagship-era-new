"""从现行v60 JSON/Markdown导出。不会运行迁移脚本或修改数据。"""
from pathlib import Path
import argparse, json, html, math, re, shutil, subprocess, tempfile, os
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[1];O=R/'outputs';O.mkdir(exist_ok=True)
d=json.loads((R/'work/v60_data.json').read_text(encoding='utf-8'))
ESC=lambda s:html.escape(str(s))
COL={'supply':'#477B87','custom':'#986047','starter':'#69747A','technology':'#236F94','contract':'#267660','facility':'#986B32','scheme':'#76578D','benchmark':'#AD4840','demand':'#526BA0'}
NAMES={'supply':'稳定供货','custom':'定制整合','starter':'初始部件','technology':'研发项目','contract':'经营合同','facility':'设施建设','scheme':'产品方案','benchmark':'参考机','demand':'市场需求'}
KIND={'chip':'芯片','screen':'屏幕','camera':'影像','body':'机身'}
market={m['id']:m['name'] for m in d['markets']}
fac={f['id']:f for f in d['factories']}

def xmltag(tag,attrs=None):
 x=OxmlElement('w:'+tag)
 for k,v in (attrs or {}).items():x.set(qn('w:'+k),str(v))
 return x

def rich(p,text):
 # 正文使用朴素Markdown，不把表格/代码标记留在Word中。
 for i,s in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)):
  r=p.add_run(s.strip('*`') if s.startswith(('**','`')) else s)
  if s.startswith('**'):r.bold=True
  if s.startswith('`'):r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)

def make_doc(mdpath,outname,title,strap):
 doc=Document();sec=doc.sections[0]
 sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.top_margin=Mm(18);sec.bottom_margin=Mm(18);sec.left_margin=sec.right_margin=Mm(18)
 sec.header_distance=Mm(8);sec.footer_distance=Mm(9)
 for name,size,bold in [('Normal',10.5,False),('Title',31,True),('Subtitle',14,False),('Heading 1',17,True),('Heading 2',12.5,True)]:
  st=doc.styles[name];st.font.name='Noto Sans CJK SC' if bold else 'Noto Serif CJK SC';st.font.size=Pt(size);st.font.bold=bold
  st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),st.font.name)
  st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.16
  st.paragraph_format.widow_control=True
  if 'Heading' in name:
   st.font.color.rgb=RGBColor.from_string('126B75');st.paragraph_format.space_before=Pt(13);st.paragraph_format.keep_with_next=True
 sec.header.paragraphs[0].text=f'旗舰元年 / {title}                                      6.0 结构重构首测版'
 for ru in sec.header.paragraphs[0].runs:ru.font.size=Pt(8);ru.font.color.rgb=RGBColor.from_string('65747B')
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 p.add_run('旗舰元年  ·  6.0  /  ').font.size=Pt(8)
 p._p.append(xmltag('fldSimple',{'instr':'PAGE'}))
 # 不涉及外部艺术资产的工业版式封面。
 doc.add_paragraph('FLAGSHIP ERA   /   6.0',style='Subtitle')
 doc.add_paragraph('旗舰元年',style='Title')
 doc.add_paragraph(title,style='Title')
 doc.add_paragraph(strap,style='Subtitle')
 doc.add_paragraph('\n')
 p=doc.add_paragraph('产品先有库存，市场才有成交。\n发布会何时发生，由玩家决定。');p.runs[0].font.size=Pt(17)
 doc.add_paragraph('\n')
 t=doc.add_table(rows=1,cols=3);t.autofit=False
 for cell,head,txt in zip(t.rows[0].cells,['两条主线','开放发布','完整内容'],['收入 / 科技\n选择较强的一条','没有固定回合数\n预热控制交易窗口','243种设计\n315张内容牌']):
  cell.width=Mm(58);cell._tc.get_or_add_tcPr().append(xmltag('shd',{'fill':'EAF1F2'}))
  p=cell.paragraphs[0];p.add_run(head).bold=True;cell.add_paragraph(txt)
 doc.add_paragraph('\n')
 doc.add_paragraph('2026年9月21日 · 测试设计，尚未完成充分多人实测。')
 doc.add_paragraph('供应资格与生产成本分开；真实库存与零售投放分开；参考机提前公开；低强度少奖，高强度多奖。')
 doc.add_page_break()
 text=mdpath.read_text(encoding='utf-8');lines=text.splitlines()
 # 前两行标题已做封面，保留后面简介与全部正文。
 i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if i<2 and line.startswith('#'):i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    vals=[s.strip() for s in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in vals):rows.append(vals)
    i+=1
   cols=len(rows[0]);table=doc.add_table(rows=1,cols=cols);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
   widths=([38,136] if cols==2 else [32,46,96] if cols==3 else [40,31,36,67] if cols==4 else [34.8]*cols)
   # 依内容给四列表格更均衡，防短数字列挤正文。
   if cols==4 and rows[0][0]=='事项':widths=[38,19,35,82]
   grid=table._tbl.tblGrid
   for child in list(grid):grid.remove(child)
   for w in widths:grid.append(xmltag('gridCol',{'w':int(w/25.4*1440)}))
   for ri,rowdata in enumerate(rows):
    row=table.rows[0] if ri==0 else table.add_row();row._tr.get_or_add_trPr().append(xmltag('cantSplit'))
    if ri==0:row._tr.get_or_add_trPr().append(xmltag('tblHeader'))
    for ci,val in enumerate(rowdata):
     cell=row.cells[ci];cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
     cell._tc.get_or_add_tcPr().append(xmltag('shd',{'fill':'DCECEF' if ri==0 else ('F3F6F7' if ri%2==0 else 'FFFFFF')}))
     mar=xmltag('tcMar')
     for ed in ['top','bottom','left','right']:mar.append(xmltag(ed,{'w':75,'type':'dxa'}))
     cell._tc.get_or_add_tcPr().append(mar)
     p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.08;rich(p,val)
     for r in p.runs:r.font.size=Pt(9.2);r.bold=ri==0
   doc.add_paragraph().paragraph_format.space_after=Pt(1)
   continue
  if line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
  elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
  elif line.startswith('# '):doc.add_paragraph(line[2:],'Heading 1')
  else:
   p=doc.add_paragraph();rich(p,line)
  i+=1
 doc.core_properties.title=f'旗舰元年6.0 {title}';doc.core_properties.author='';doc.core_properties.subject='产品平台、收入与科技、玩家推动发布、随机参考机'
 doc.save(O/outname)

CSS='''
*{box-sizing:border-box}body{margin:0;font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;color:#24343B;background:#e6ebee}h1,h2,h3,p{margin:0}button{font:inherit;cursor:pointer;border:1px solid #7898A2;background:white;padding:8px 12px;border-radius:4px;color:#234550}header.web{background:#18323E;color:white;padding:16px 24px;display:flex;gap:16px;align-items:center;flex-wrap:wrap}header.web h1{font-size:23px}.page{width:194mm;height:281mm;background:white;margin:15px auto;position:relative;padding:0;overflow:hidden}.pagehead{height:12mm;border-bottom:1px solid #B2C8D0;display:flex;align-items:center;justify-content:space-between;font-size:10pt;color:#315962}.pagefoot{position:absolute;bottom:0;left:0;right:0;border-top:1px solid #DAE3E8;padding-top:2mm;font-size:7.5pt;color:#667984;display:flex;justify-content:space-between}.cardgrid{display:grid;grid-template-columns:repeat(2,1fr);grid-template-rows:repeat(3,84mm);gap:3mm;margin-top:4mm}.cardgrid.tall{grid-template-rows:repeat(2,125mm)}.card{position:relative;border:1px solid #B8C8CF;border-top:4px solid var(--accent);padding:3.4mm 3.2mm 7mm;background:white;overflow:hidden}.eyebrow{font-size:7.3pt;letter-spacing:.3px;display:flex;justify-content:space-between;color:var(--accent);margin-bottom:1.2mm}.card h3{font-size:12pt;line-height:1.22;margin-bottom:2mm;letter-spacing:.1px}.metrics{display:flex;flex-wrap:wrap;gap:1mm;margin:2mm 0}.metric{background:#EDF2F4;padding:1mm 1.5mm;font-size:8pt;line-height:1.2}.metric b{font-size:11.5pt;margin-left:1mm}.card p{font-size:8.7pt;line-height:1.48;margin:1.8mm 0;overflow-wrap:anywhere}.card p.small{font-size:7.4pt;line-height:1.4;color:#536873}.effect{border-top:1px solid #CCDADD;padding-top:2mm}.card .bottom{position:absolute;bottom:2mm;left:3.2mm;right:3.2mm;font-size:7pt;color:#6A7C83;display:flex;justify-content:space-between}.card table{border-collapse:collapse;width:100%;font-size:8pt;line-height:1.35;margin:1.5mm 0}.card th{background:#E9F0F3;font-weight:500}.card td,.card th{padding:1mm .8mm;text-align:center;border-bottom:1px solid #DCE6EA}.card td:first-child{text-align:left}.label{font-size:7.5pt;font-weight:bold;color:#47656F}.factory{display:inline-block;font-size:8pt;padding:.8mm 1.5mm;border:1px solid var(--fc);color:var(--fc);border-radius:3px;margin:1mm 0}.cover{padding:15mm 9mm;background:linear-gradient(145deg,#F5F8F9,#E4EDF1)}.cover .big{font-size:38pt;font-weight:900;color:#183A47;margin-top:22mm}.cover .sub{font-size:20pt;color:#246D79;margin-top:8mm}.cover .lede{font-size:13pt;line-height:1.9;margin-top:20mm;max-width:150mm}.cover .statrow{display:flex;gap:15mm;margin-top:22mm;color:#153D4A}.cover .statrow strong{font-size:30pt;display:block}.index{padding:8mm 5mm;font-size:10pt}.index h2{font-size:20pt;margin-bottom:7mm;color:#176E79}.index p{font-size:10pt;line-height:1.8;margin:3mm 0}.index table{border-collapse:collapse;width:100%;font-size:10pt;margin:6mm 0}.index td,.index th{text-align:left;padding:3mm;border-bottom:1px solid #CDDBE0}.index th{background:#E9F2F3}.index .note{background:#E9F1F3;padding:5mm;line-height:1.8;margin:5mm 0}.cover .watermark{font-size:9pt;color:#506973;margin-top:25mm}@page{size:A4 portrait;margin:8mm}@media print{body{background:white}header.web{display:none}.page{margin:0;page-break-after:always;break-after:page}.page:last-child{page-break-after:auto;break-after:auto}*{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
'''

def factory_html(id):
 f=fac[id];return f'<span class="factory" style="--fc:{f["color"]}">{ESC(f["name"])} · {ESC(f["shape"])}</span>'

def card_body(c):
 t=c['type'];out=[]
 if t in ('supply','custom','starter'):
  out.append(f'<div class="metrics"><span class="metric">供货代 <b>{"初始" if t=="starter" else "I II III IV V".split()[c["era"]-1]}</b></span><span class="metric">整卡含<b>{len(c["components"])}项</b></span></div>')
  rows='<tr><th>部件</th><th>成本</th><th>规格</th><th>耗/供</th><th>标识</th></tr>'
  for p in c['components']:
   rows+=f'<tr><td>{KIND[p["kind"]]}<br>{"折叠" if p["form"]=="fold" else "直板"}</td><td>{p["batchCost"]}</td><td>{p["spec"]}</td><td>{"供" if p["kind"]=="body" else "耗"}{p["power"]}</td><td>{ESC(p["mark"])}</td></tr>'
  out.append('<table>'+rows+'</table>')
  out.append(f'<p class="effect">{ESC(c["effect"])}</p>')
  out.append('<p class="small">型号退出前持续占用。稳定/初始退市回库，定制退市移出。取得费与每批生产费分开。</p>')
 elif t=='technology':
  out.append('<div class="metrics">'+''.join(f'<span class="metric">{label}<b>{c[key]}</b></span>' for label,key in [('费用','cost'),('工期','work'),('科技','gain')])+'</div>')
  out.append(f'<p class="small">领域：{ESC(c["field"])}；此前完成项目≥{c["prerequisiteProjects"]}</p>')
  out.append(f'<p><span class="label">完成证据　</span>{ESC(c["proof"])}</p>')
  out.append(f'<p class="effect">{ESC(c["effect"])}</p>')
 elif t=='contract':
  out.append(factory_html(c['factory']))
  out.append(f'<div class="metrics"><span class="metric">期限<b>{c["periods"]}次发布</b></span><span class="metric">{"押金" if c.get("deposit",0) else "费用"}<b>{c.get("deposit") or c.get("fee",0)}</b></span></div>')
  if c['subtype']=='order':out.append(f'<p><span class="label">交付 {c["units"]}批　</span>每批{c["unitPrice"]}现金 / {c["unitIncome"]}收入</p>')
  out.append(f'<p class="small"><span class="label">条件　</span>{ESC(c["requirement"])}</p><p class="effect">{ESC(c["effect"])}</p>')
 elif t=='facility':
  out.append(f'<div class="metrics"><span class="metric">建设<b>{c["cost"]}</b></span><span class="metric">工作<b>{c["work"]}</b></span><span class="metric">空间<b>{c["space"]}</b></span></div><p class="effect">{ESC(c["effect"])}</p><p class="small">先取入手，再付工作与现金建设。不提供厂家积累，不附带科技分。</p>')
 elif t=='scheme':
  out.append(f'<div class="metrics"><span class="metric">额外开发费<b>{c["developmentCost"]}</b></span><span class="metric">每型号最多<b>1方案</b></span></div><p><span class="label">条件　</span>{ESC(c["requirement"])}</p><p class="effect">{ESC(c["effect"])}</p><p class="small">可选设计：没有Q牌也能造手机。方案不印上市收入；实际成交才计收入。</p>')
 elif t=='benchmark':
  out.append(f'<div class="metrics"><span class="metric">参考机世代<b>{c["era"]}</b></span><span class="metric">每市场容量<b>1</b></span></div>')
  rows='<tr><th>市场</th><th>竞争力</th><th>挑战奖励（二选一）</th></tr>'
  for m in c['markets']:
   rw='无' if m['reward']==0 else f'{m["reward"]*2}收入 / {m["reward"]}科技'
   rows+=f'<tr><td>{market[m["market"]]}</td><td><b>{m["score"]}</b></td><td>{rw}</td></tr>'
  out.append('<table>'+rows+'</table><p class="small">'+ESC(c['effect'])+'</p>')
 elif t=='demand':
  rows='<tr><th>市场</th><th>预算</th><th>订单量</th><th>偏好</th></tr>'
  for m in c['markets']:
   delta=m['ordersDelta'];s='人数' if not delta else f'人数{delta:+}'
   rows+=f'<tr><td>{market[m["market"]]}</td><td>{m["budget"]}</td><td>{s}</td><td>{" / ".join(m["preferences"])}</td></tr>'
  out.append('<table>'+rows+'</table><p class="effect">'+ESC(c['effect'])+'</p><p class="small">预算不等于售价。每个偏好通常只加一次。实际发布后才换需求，没有固定轮数。</p>')
 return ''.join(out)

def render_card(c):
 foot='每期芯片平台通常2批' if c['type'] in ('supply','starter','custom') else ('工艺 / 公司' if c.get('scope')=='process' else NAMES[c['type']])
 return f'<article class="card" id="{c["id"]}" style="--accent:{COL[c["type"]]}"><div class="eyebrow"><span>{c["id"]} / {NAMES[c["type"]]}</span><span>实体 ×{c["qty"]}</span></div><h3>{ESC(c["name"])}</h3>{card_body(c)}<div class="bottom"><span>{foot}</span><span>旗舰元年 6.0</span></div></article>'

def build_catalog():
 groups=[('稳定供货', [c for c in d['parts'] if c['type']=='supply'],6),('定制整合',[c for c in d['parts'] if c['type']=='custom'],4),('初始部件',[c for c in d['parts'] if c['type']=='starter'],6),('研发项目',d['technologies'],6),('经营合同',d['contracts'],6),('设施建设',d['facilities'],6),('产品方案',d['schemes'],6),('随机参考机',d['benchmarks'],6),('公开需求',d['demands'],4)]
 pages=[];manifest=[];ranges=[];page=3
 for title,cards,n in groups:
  start=page
  for j in range(0,len(cards),n):
   batch=cards[j:j+n]
   cardhtml=''.join(render_card(c) for c in batch)
   pages.append(f'<section class="page"><div class="pagehead"><b>{title}</b><span>6.0 / 完整图鉴</span></div><div class="cardgrid {"tall" if n==4 else ""}">{cardhtml}</div><div class="pagefoot"><span>测试数值 · 不能与5.x混用 · 不是量产裁切稿</span><span>{page:02}</span></div></section>')
   for c in batch:manifest.append({'id':c['id'],'name':c['name'],'type':c['type'],'page':page,'qty':c['qty']})
   page+=1
  ranges.append((title,len(cards),start,page-1))
 cover='<section class="page cover"><p>FLAGSHIP ERA / 6.0</p><div class="big">旗舰元年</div><div class="sub">完整卡牌图鉴</div><div class="lede">以真实产品、研发项目、限期合同与产业设施构筑公司。参考机按世代随机出现，准备期即可看清对手。</div><div class="statrow"><div><strong>243</strong>种设计</div><div><strong>315</strong>张内容牌</div></div><p class="watermark">2026年9月21日 · 结构重构首测版<br>所有卡面完整列出；重复副本按×份数准备。<br>不含通用卡背、出血及实际印厂刀模。</p></section>'
 index='<section class="page index"><h2>索引与读卡方式</h2><p>本图鉴与6.0完整规则共同使用。不能因编号沿用5.3，就继续用5.3的费用、门槛或生命周期。</p><table><tr><th>内容</th><th>设计数</th><th>页码</th></tr>'+''.join(f'<tr><td>{a}</td><td>{b}</td><td>{c}—{e}</td></tr>' for a,b,c,e in ranges)+'</table><div class="note">研发的“工期”用工作点实际推进；合同期限按真实发布次数减少，不是按全桌轮次。<br>部件“成本”为每批生产成本，取得方案仍另支付签约或定制费用。<br>部件上的标识不等于整机已取得六种特征；按规则由规格、结构、供电或明确能力推导。</div><p>本次没有额外随机事件罚人，也没有强制企划卡。产品方案Q是可选的差异化设计。参考V一张覆盖四市场，挑战奖励必须按实际击败的那一格读取。</p><div class="pagefoot"><span>规则与规划分册，不用查旧版补丁</span><span>02</span></div></section>'
 h='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>旗舰元年6.0 全卡牌图鉴</title><style>'+CSS+'</style><body><header class="web"><h1>旗舰元年 6.0 · 全卡牌图鉴</h1><button onclick="window.print()">打印图鉴</button><span>静态查阅文件 · 243种 / 315张</span></header>'+cover+index+''.join(pages)+'</body></html>'
 (O/'旗舰元年_6.0_全卡牌图鉴.html').write_text(h,encoding='utf-8')
 (R/'work/catalog_manifest.json').write_text(json.dumps({'cards':manifest,'pages':page-1,'ranges':ranges},ensure_ascii=False,indent=2),encoding='utf-8')
 # 完整可编辑文字图鉴，所有字段原样明确列出。
 text=['# 旗舰元年6.0 完整文字图鉴','243种设计 / 315张实体内容牌。函数生成，数据为work/v60_data.json。']
 for title,cards,n in groups:
  text+=['',f'## {title}']
  for c in cards:
   text+=['',f'### {c["id"]} {c["name"]}（×{c["qty"]}）']
   for k,v in c.items():
    if k in ['id','name','qty','type']:continue
    text.append(f'{k}：'+(json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else str(v)))
 (R/'work/v60_catalog.md').write_text('\n\n'.join(text),encoding='utf-8')

BOARD_CSS='''
*{box-sizing:border-box}body{margin:0;font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;background:#E7ECEF;color:#233C47}header{background:#173442;color:white;padding:18px;display:flex;gap:8px;flex-wrap:wrap;align-items:center}h1{font-size:23px;margin:0 22px 0 0}button{padding:8px 10px;border:1px solid #6E959F;background:white;color:#234551;cursor:pointer;font:12px "Noto Sans CJK SC"}.board{width:400mm;height:277mm;background:#F9FBFC;margin:16px auto;position:relative;padding:9mm;overflow:hidden;display:none}.board.active{display:block}.title{border-bottom:2px solid #198091;padding-bottom:5mm;display:flex;justify-content:space-between;align-items:end;margin-bottom:6mm}.title h2{font-size:25pt;margin:0}.title p{font-size:10pt;margin:0}.subtitle{font-size:11pt;line-height:1.7;color:#58717A;margin:3mm 0 5mm}.row{display:flex;gap:4mm}.col{flex:1}.panel{background:white;border:1px solid #ACBEC7;padding:4mm;margin-bottom:4mm}.panel h3{margin:0 0 3mm;font-size:15pt;color:#146E7B}.panel p,.panel li{font-size:10pt;line-height:1.6;margin:2mm 0}.slots{display:flex;gap:3mm}.slot{flex:1;border:1px dashed #8BA6B1;background:#F1F6F8;min-height:25mm;padding:3mm;text-align:center;font-size:11pt;color:#506C78}.slot b{font-size:17pt;display:block}.track{display:flex;gap:1mm;flex-wrap:wrap}.tick{min-width:11mm;height:12mm;padding:1mm;border:1px solid #A9BDC7;background:white;text-align:center;font-size:11pt;display:inline-flex;align-items:center;justify-content:center;flex-direction:column}.tick small{font-size:6.5pt}.tick.end{border:2px solid #C25E30;background:#FFF1DE}.grid2{display:grid;grid-template-columns:1fr 1fr;gap:4mm}.grid3{display:grid;grid-template-columns:repeat(3,1fr);gap:4mm}.legend{position:absolute;bottom:6mm;left:9mm;right:9mm;border-top:1px solid #ADC4CD;padding-top:2mm;font-size:8.5pt;color:#54707C}table{border-collapse:collapse;width:100%;font-size:10pt}td,th{padding:2mm;border-bottom:1px solid #D5E0E5;text-align:left}th{background:#EAF2F5}.money{color:#9A6326;font-weight:bold}.focus{font-size:19pt;color:#197586;margin:3mm 0;font-weight:bold}.light{background:#EAF3F6}.materialslots .slot{height:81mm}.modelmeta .tick{height:9mm;min-width:8mm}.offerrow{display:flex;align-items:center;justify-content:space-between;gap:2mm;font-size:10pt;margin:2mm 0}.offerrow .tick{height:9mm;min-width:9mm}.scores{display:grid;grid-template-columns:repeat(17,1fr);gap:1mm}.scores .tick{width:auto;min-width:0;height:10mm}.techscores{display:grid;grid-template-columns:repeat(19,1fr);gap:1mm}.techscores .tick{height:17mm;min-width:0;width:auto}.techscores .tick b{font-size:13pt}.footertext{font-size:10pt;line-height:1.7}#market .panel,#company .panel{padding:3mm}#market .panel p,#company .panel p{font-size:9pt}#market .slot{min-height:20mm}#supply .slot{min-height:20mm!important}#supply .panel{padding:3mm;margin-bottom:3mm}#supply .panel p{font-size:9pt;margin:1mm 0}#offers .slot{height:40mm!important;min-height:0}#company .slot{height:30mm!important;min-height:20mm}@page{size:A3 landscape;margin:10mm}@media print{header{display:none}body{background:white}.board{display:none;margin:0;break-after:page}.board.active,body.printall .board{display:block}.board:last-child{break-after:auto}*{print-color-adjust:exact;-webkit-print-color-adjust:exact}}
'''

def ticks(n,start=0,ends=(),small=False):return '<div class="track">'+''.join(f'<span class="tick {"end" if i in ends else ""}">{i}</span>' for i in range(start,n+1))+'</div>'
def panel(h,content,cls=''):return f'<div class="panel {cls}"><h3>{h}</h3>{content}</div>'
def slots(labels,height=''):return '<div class="slots">'+''.join(f'<div class="slot" style="{height}">{x}</div>' for x in labels)+'</div>'
def board(id,title,sub,body):return f'<section class="board" id="{id}"><div class="title"><h2>{title}</h2><p>旗舰元年 / 6.0</p></div><p class="subtitle">{sub}</p>{body}<div class="legend">仅展示、切换与打印 · 不发牌、不计算、不保存棋局 · 全部数字为首测参数 · 纸牌可置于对应区域旁</div></section>'

def build_boards():
 boards=[]
 body=panel('共同发布',ticks(20,ends=(12,16,20))+'<p>2人12 / 3人16 / 4人20。每本人回合可花1工作点推进1—3格。到顶在当前玩家工作结束后发布，没有其他玩家响应回合。</p>')
 body+='<div class="row">'+panel('本场公开环境',slots(['需求N<br>提前公开','参考V<br>按代随机抽1']))+panel('所有型号都能投放','<p>先分合同与零售库存。各市场按竞争力从高到低分配有限订单。不是只取每家公司最强一台。</p><p>参考每市场最多占1份；同分真人优先；严格超过且真实零售成交才有挑战资格。</p>')+'</div>'
 body+='<div class="grid2">'+''.join(panel(m['name'],f'<p>{m["eligibility"]}</p>'+slots(['订单量<br>人数＋修正','预算 / 偏好<br>按当前N'])+'<p>此处放已锁定的型号投放标识；库存不得重复承诺。</p>') for m in d['markets'])+'</div>'
 boards.append(board('market','开放市场与共同发布','收入主要来自售出的实际批次；库存没卖出就留下，不按上市卡片发固定收入。',body))
 body=''
 for k,n in KIND.items():
  body+=panel(n+'公开供货',slots([f'<b>{i}</b>'+('最旧' if i==1 else '新') for i in range(1,6)],'min-height:29mm')+'<p>展示人数＋1张，取走不补；不足展示实际张数。完整牌可置于本行旁。</p>')
 boards.append(board('supply','分代供应链','当前代签约9，早一代6，早至少两代3；只开放当前及过去世代。芯片平台通常每周期支持2批。',body))
 body=panel('研发项目 / 4张公开',slots([f'研发 {i}' for i in range(1,5)],'height:69mm')+'<p>可1工作点取入手，或立项时从这里直接取得并付款、放1进度。位置无强度限制。</p>')
 body+=panel('商业牌 / 6张公开',slots([f'商业 {i}' for i in range(1,7)],'height:48mm')+'<p>1工作点至多取2张。合同H、设施U、产品方案Q在同一商业牌堆；先取牌，再签约或建设。</p>')
 body+=panel('定制整合 / 3张公开',slots(['定制1','定制2','定制3'])+'<p>1工作点取得1张，付6；全部部件必须同型号，不拆分。制造成本另付。</p>')
 boards.append(board('offers','技术与商业机会','公开牌取走立即补；稳定供货行例外不立即补。未开放定制不能提前进入本区。',body))
 body=panel('本回合工作',slots(['<b>1</b>工作','<b>2</b>工作','<b>3</b>工作'])+'<p>每个自己的回合重新取得3点，不保存；没有五行动强度或行动卡冷却。</p>')
 body+='<div class="grid2">'+panel('在研 / 默认2槽',slots(['项目A<br>工期0—4','项目B<br>工期0—4'],'height:43mm'))+panel('执行合同 / 默认2槽',slots(['合同A<br>还可经历2 / 1次发布','合同B<br>还可经历2 / 1次发布'],'height:43mm'))+'</div>'
 body+=panel('设施 / 默认3格',slots(['设施空间1','设施空间2','设施空间3'],'height:33mm')+'<p>设施卡占1或2格。建设通常同回合2工作点，按卡面付现金；不自动增加收入科技。</p>')
 body+=panel('本期生产产能',ticks(10)+'<p>通常上限6，设施最高10；旧机补产与新机备货共同计数。只有真实发布后复位。</p>')
 boards.append(board('company','公司工程排期','工作点每本人回合补充；产能、芯片配额、合同期限依真实发布处理。这不是固定轮次游戏。',body))
 body='<div class="row modelmeta">'+panel('型号编号',ticks(4,start=1))+panel('年龄',ticks(3)+'<p>0首发 → 1 → 2后退市；T16例外至3。</p>')+panel('库存',ticks(8)+'<p>默认最多4，扩容最高8。</p>')+'</div>'
 body+='<div class="materialslots">'+slots(['<b>芯片</b>附本期平台用量<br>或T37共享占位','<b>屏幕</b>折叠必须配套','<b>影像</b>高级模组检查兼容','<b>机身</b>提供供电'])+'</div>'
 body+='<div class="row" style="margin-top:4mm">'+panel('部署技术 / 通常2槽',slots(['技术1','技术2','T21额外槽'],'height:42mm'))+panel('可选Q方案',slots(['每型号最多1张<br>没有Q也可定型'],'height:42mm'))+'</div>'
 body+='<p class="footertext">标记六特征：性能 / 影像 / 轻薄 / 续航 / 折叠 / 生态。基础竞争力不含价格、偏好、市场特效与年龄。<br>型号验证：未验证 / 已验证。本型号的部件和技术持续占用直到合法改款或退市。</p>'
 boards.append(board('model','型号开发面板','每人重复准备四份，分别编号1—4。新型号进屏风，公开后移到在售区；不用手写配置。',body))
 body='<div class="grid2">'
 for i in range(1,5):
  rows=''.join('<div class="offerrow"><span>'+m['name']+'</span>'+ticks(4)+'</div>' for m in d['markets'])
  body+=panel(f'型号{i}的零售投放',rows+'<p>通常每市场至多3批。合同留货先扣；四市场投放合计不能超过剩余库存。</p>')
 body+='</div>'+panel('价格一起扣下','<p>同一型号在所有零售市场使用同一价格。18 / 28 / 38 / 48。48需通常旗舰条件或明确许可。数量与价格均先锁定后揭晓。</p>')
 boards.append(board('allocation','秘密库存分配','每个数量行放1个小方块即可；0为不投放。此表放屏风内，锁定后与价格一起揭晓。',body))
 body=panel('收入 / 商业主线','<div class="scores">'+''.join(f'<span class="tick {"end" if i==60 else ""}">{i}{"<small>终局门槛</small>" if i==60 else ""}</span>' for i in range(101))+'</div>')
 body+=panel('科技 / 研发主线','<div class="techscores">'+''.join(f'<span class="tick {"end" if i==30 else ""}"><b>{i}</b><small>成绩 {i*2}</small></span>' for i in range(37))+'</div>')
 body+=panel('任一达标，不把两项相加','<div class="focus">收入 ≥ 60　或　科技 ≥ 30</div><p>完成当前回合全部结算后，所有玩家各再一回合，再共同最终发布。企业成绩取收入与2×科技的较大者。平分先比另一项，再比现金。</p>')
 boards.append(board('scores','两条可选择的主线','两条记录分别增长，不交会。不要求收入和科技同时发展，没有秘密目标附加分。',body))
 body='<div class="grid2">'+panel('竞品挑战奖励','<table><tr><th>参考格数值</th><th>星级</th><th>整份二选一</th></tr><tr><td>0—10</td><td>0</td><td>无</td></tr><tr><td>11—16</td><td>1</td><td>2收入 / 1科技</td></tr><tr><td>17—24</td><td>2</td><td>4收入 / 2科技</td></tr><tr><td>25+</td><td>3</td><td>6收入 / 3科技</td></tr></table><p>必须严格胜出且真实零售成交。每公司每场最多一次。科技奖还须首发＋部署研发。</p>')+panel('真实售价与每批收入','<table><tr><th>售价</th><th>竞争修正</th><th>收入/批</th></tr><tr><td>18</td><td>+6</td><td>1</td></tr><tr><td>28</td><td>+2</td><td>2</td></tr><tr><td>38</td><td>−2</td><td>3</td></tr><tr><td>48</td><td>−6</td><td>4</td></tr></table><p>高于预算不成交。按真实批数领取现金，未售库存留存。不是固定第一/第二收入表。</p>')+'</div>'
 body+=panel('发布时间与产品寿命','<p>准备 → 玩家预热到顶 → 当前回合结束 → 同时锁定揭晓 → 合同与有限零售 → 收款/成果 → 老化/退市 → 抽新需求和参考。</p><p>首发年龄0；下一场年龄1减3；再下一场年龄2减6，然后退市。没有每全桌一轮自动结算，也没有过回合自动老化。</p>')
 body+=panel('参考机：同代有差异','<p>每代6款随机抽1张全桌共用。弱市场低奖，强市场高奖，不跨格套奖励。当前准备期即公开，中途解锁新供货不改当前参考，下一场才抽新代。</p>')
 boards.append(board('reference','价格、参考机与生命周期','虚拟参考每市场只占至多1份订单，不是自动行动对手。实际批次数与排名才决定销售。',body))
 htmls='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>旗舰元年6.0 线下版图</title><style>'+BOARD_CSS+'</style><body><header><h1>旗舰元年 6.0</h1>'+''.join(f'<button data-id="{re.search(r"id=\"([^\"]+)",b).group(1)}">{n}</button>' for b,n in zip(boards,['市场与发布','供应链','牌市','公司','型号','秘密投放','两条主线','参考速查']))+'<button onclick="window.print()">打印当前</button><button onclick="document.body.classList.add(\'printall\');window.print();document.body.classList.remove(\'printall\')">打印全部</button></header>'+''.join(boards)+'''<script>document.querySelector('.board').classList.add('active');document.querySelectorAll('button[data-id]').forEach(b=>b.onclick=()=>{document.querySelectorAll('.board').forEach(p=>p.classList.toggle('active',p.id===b.dataset.id))});</script></body></html>'''
 (O/'旗舰元年_6.0_线下版图.html').write_text(htmls,encoding='utf-8')
 (R/'work/boards_manifest.json').write_text(json.dumps({'views':8,'ids':['market','supply','offers','company','model','allocation','scores','reference'],'onlyDisplayAndPrint':True},ensure_ascii=False,indent=2))

AID_CSS='''*{box-sizing:border-box}body{margin:0;background:#E8EFF2;font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;color:#213E4B}.sheet{width:194mm;height:281mm;background:white;margin:12px auto;position:relative;overflow:hidden}.top{height:12mm;border-bottom:1px solid #95B4C1;display:flex;justify-content:space-between;align-items:center;font-size:10pt}.foot{position:absolute;bottom:0;left:0;right:0;border-top:1px solid #CCDADF;font-size:7.5pt;padding-top:2mm;color:#54707B}.pricegrid{display:grid;grid-template-columns:repeat(4,46mm);grid-template-rows:repeat(4,62mm);gap:3mm;width:193mm;margin:4mm auto 0}.price{border:1px dashed #8FA9B5;text-align:center;position:relative;padding:4mm 2mm;color:#27515F;overflow:hidden}.price .n{font-size:10pt}.price .v{font-size:30pt;margin:4mm 0 1mm;font-weight:900;line-height:1.1}.price p{font-size:8.5pt;line-height:1.6;margin:2mm 0}.price.back{background:repeating-linear-gradient(45deg,#EAF0F3,#EAF0F3 3mm,#F5F8FA 3mm,#F5F8FA 6mm);border:1px dashed #8FA9B5}.price.back .v{margin-top:13mm}.panel{border:1px solid #B8CBD3;padding:3mm;margin:3mm 0}.panel h3{margin:0 0 2mm;font-size:12pt;color:#206F7A}.panel p{font-size:8.5pt;line-height:1.6;margin:2mm 0}.track{display:flex;gap:1.5mm;align-items:center}.tick{display:inline-flex;align-items:center;justify-content:center;width:8mm;height:8mm;border:1px solid #A5BDC7;font-size:9pt}.orow{display:flex;justify-content:space-between;margin:2mm 0;align-items:center;font-size:9pt}.cols{display:grid;grid-template-columns:1fr 1fr;gap:3mm}.tokens{display:grid;grid-template-columns:repeat(6,1fr);gap:2mm;margin-top:3mm}.token{border:1px dashed #9FB5BF;min-height:20mm;text-align:center;font-size:9pt;padding:3mm 1mm;line-height:1.5}.info{font-size:10pt;line-height:1.9;margin:7mm 2mm}table{width:100%;border-collapse:collapse;font-size:10pt}th,td{padding:3mm;border-bottom:1px solid #D1DDE3;text-align:left}th{background:#ECF4F6}h2{font-size:18pt;color:#176F7E}header{padding:12px 20px;background:#173B4D;color:white;font-size:18px}button{padding:7px 12px;margin-left:15px}@page{size:A4 portrait;margin:8mm}@media print{body{background:white}header{display:none}.sheet{margin:0;break-after:page}.sheet:last-child{break-after:auto}*{print-color-adjust:exact;-webkit-print-color-adjust:exact}}'''

def aid_sheet(title,body,n,note):return f'<section class="sheet"><div class="top"><b>{title}</b><span>6.0 实体组件 / {n}</span></div>{body}<div class="foot">{note}</div></section>'
def atrack(n):return '<div class="track">'+''.join(f'<span class="tick">{i}</span>' for i in range(n+1))+'</div>'

def build_aids():
 sheets=[];prices=d['parameters']['prices'];mods=d['parameters']['priceScore'];inc=d['parameters']['incomePerBatch']
 front=[];back=[]
 for model in range(1,5):
  for p,m,i in zip(prices,mods,inc):
   front.append(f'<div class="price"><div class="n">型号 {model} / 售价</div><div class="v">{p}</div><p>竞争修正 {m:+}<br>成交每批：{p}现金<br>收入 +{i}</p><p>{"需满足48售价许可" if p==48 else "仍须符合顾客预算"}</p></div>')
  for p in reversed(prices):back.append(f'<div class="price back"><div class="n">旗舰元年 6.0</div><div class="v">{model}</div><p>型号秘密定价<br>正面仅本人可见</p></div>')
 sheets.append(aid_sheet('16张价格牌 / 正面','<div class="pricegrid">'+''.join(front)+'</div>',1,'每人一套。第1—2页A4原尺寸、长边翻转双面。先试印，实际打印机套准未验证。'))
 sheets.append(aid_sheet('16张价格牌 / 背面','<div class="pricegrid">'+''.join(back)+'</div>',2,'背面同型号完全相同，不透露18/28/38/48。裁切可分别贴合；无需写字。'))
 body='<div class="cols">'
 for model in range(1,5):
  rows=''.join('<div class="orow"><span>'+m['name']+'</span>'+atrack(4)+'</div>' for m in d['markets'])
  body+=f'<div class="panel"><h3>型号 {model} 投放</h3>{rows}<p>通常每行至多3；特许可4。<br>先预留合同库存，合计不得超库存。</p></div>'
 body+='</div><div class="panel"><h3>锁定顺序</h3><p>先给订单合同分货，再给各零售市场分货，剩余留库。价格同型号所有市场一致。全部玩家锁定后一起揭晓。</p><p>每行放1个小方块在数字格；无库存、未获得市场资格或高于预算，都不能成交。揭晓后不转移失败市场的库存。</p></div>'
 sheets.append(aid_sheet('秘密投放垫板',body,3,'每人1份单面，放屏风内。准备16个数量小方块；纸片可用第6页通用标记代替。'))
 body=''
 for model in range(1,5):
  body+=f'<div class="panel"><h3>型号 {model} 状态</h3><div class="orow"><span>库存</span>{atrack(8)}</div><div class="orow"><span>年龄</span>{atrack(3)}</div><div class="orow"><span>独立验证</span><span>□ 未做　□ 已做</span></div><p>0首发，1/2持续销售，年龄2发布后退市；T16例外。改款不清除年龄/验证。</p></div>'
 sheets.append(aid_sheet('四款型号状态',body,4,'每人1份单面；库存通常上限4，明确能力提高后至多8。方框放片，不用勾写。'))
 body=panel('本期公司产能',atrack(10)+'<p>通常6，最高10；每次生产累加，实际发布结束才复位。</p>')
 for i in range(10):body+=f'<div class="panel" style="padding:2mm;margin:2mm 0"><div class="orow"><b style="font-size:9pt">平台用量</b>{atrack(5)}</div></div>'
 sheets.append(aid_sheet('芯片平台配额条',body,5,'每人一页单面。剪下额度条随实际芯片放置；退市回库也保留本期用量，防止重新定型刷新配额。'))
 toks=['工作1','工作2','工作3','本期验证\n可用','本期验证\n已用','本场挑战\n已领取']
 toks+=['合同期限2','合同期限1','合同期限0','进度0','进度1','进度2','进度3','进度4','设施停用','实验室加速\n本回合已用']
 toks += [f'共享 {a}→{b}' for a in range(1,5) for b in range(1,5) if a!=b]
 toks += ['性能','影像','轻薄','续航','折叠','生态','安全资格','耐候资格','收入+100','科技+36']
 toks += ['数量片']*24
 body='<div class="tokens">'+''.join(f'<div class="token">{ESC(x).replace(chr(10),"<br>")}</div>' for x in toks)+'</div>'
 sheets.append(aid_sheet('通用指示片与共享占位',body,6,'每人按需打印；可用现有小方块替代数量与状态片。同批库存不得在合同、多个零售或回收中重复使用。'))
 body='<div class="info"><h2>每个人回合：3工作点</h2><p>预热花1点，推进1—3；每本人回合一次。到顶在当前回合结束立即发布。不是全员行动一轮就结算。</p></div>'
 body+='<table><tr><th>参考格竞争力</th><th>星级</th><th>一次奖二选一</th></tr><tr><td>0—10</td><td>0</td><td>无</td></tr><tr><td>11—16</td><td>1</td><td>2收入 / 1科技</td></tr><tr><td>17—24</td><td>2</td><td>4收入 / 2科技</td></tr><tr><td>25以上</td><td>3</td><td>6收入 / 3科技</td></tr></table>'
 body+='<div class="info"><p>真实零售至少成交1批，竞争力严格超过那一市场参考值；每公司每场最多领取一次。<b>选科技还须首发并部署完成研发。</b>同分先成交但不能领奖；不能拿弱项的胜出领强项高奖。</p><h2>两条可选主线</h2><p>收入60或科技30，任一达标触发最后阶段；所有玩家各再一回合，然后最终发布。<br>企业成绩：收入与科技×2取较大者；不是相加或标记相遇。</p><h2>印刷与准备</h2><p>价格第1—2页双面，其余单面。每人状态面板与型号面板按需要重复准备。新型号四部件与部署技术保留到退市，不用手写旧配置。</p></div>'
 sheets.append(aid_sheet('桌边规则速查',body,7,'全桌1份。参考机和需求实际卡面见完整图鉴V/N；本页不是可随机抽取的参考机。'))
 h='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>旗舰元年6.0 实体组件</title><style>'+AID_CSS+'</style><body><header>旗舰元年6.0 · 实体组件<button onclick="window.print()">打印</button></header>'+''.join(sheets)+'</body></html>'
 (O/'旗舰元年_6.0_实体组件.html').write_text(h,encoding='utf-8')
 (R/'work/aids_manifest.json').write_text(json.dumps({'pages':7,'pricesPerPlayer':16,'priceSizeMM':[46,62],'duplexPages':[1,2],'singleSidedPages':[3,4,5,6,7],'handwritingRequired':False},ensure_ascii=False,indent=2))

def export_pdf():
 # 安装系统LibreOffice，以及Chromium或playwright chromium后可独立运行。
 lo=shutil.which('libreoffice') or shutil.which('soffice')
 if not lo:raise RuntimeError('需要LibreOffice/soffice才能导出Word对应PDF')
 for p in O.glob('*.docx'):
  profile=Path(tempfile.mkdtemp(prefix='flagship60-lo-')).as_uri()
  subprocess.run([lo,f'-env:UserInstallation={profile}','--headless','--convert-to','pdf','--outdir',str(O),str(p)],check=True,timeout=120)
 from playwright.sync_api import sync_playwright
 chrome=os.environ.get('FLAGSHIP_CHROMIUM') or shutil.which('chromium') or shutil.which('chromium-browser')
 with sync_playwright() as pw:
  args={'headless':True,'args':['--no-sandbox']}
  if chrome:args['executable_path']=chrome
  browser=pw.chromium.launch(**args)
  for p in O.glob('*.html'):
   page=browser.new_page(viewport={'width':1600,'height':1200})
   page.set_content(p.read_text(encoding='utf-8'),wait_until='load');page.evaluate('document.fonts.ready')
   if '线下版图' in p.name:page.evaluate("document.body.classList.add('printall')")
   page.pdf(path=str(p.with_suffix('.pdf')),prefer_css_page_size=True,print_background=True)
   page.close()
  browser.close()

if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--no-pdf',action='store_true');a=ap.parse_args()
 make_doc(R/'work/v60_rules.md','旗舰元年_6.0_完整规则书.docx','完整规则书','产品线 · 批次生产 · 玩家推动发布')
 make_doc(R/'work/v60_planning.md','旗舰元年_6.0_设计规划书.docx','设计规划书','结构重构 · 两条主线 · 随机竞品')
 build_catalog();build_boards();build_aids()
 if not a.no_pdf:export_pdf()
 print('Built current 6.0 files. Migration was not run.')
